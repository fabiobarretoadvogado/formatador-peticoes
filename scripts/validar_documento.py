from __future__ import annotations

import argparse
import re
from collections import Counter
from pathlib import Path

from docx import Document

try:
    from .classificar_regras import RE_ASSINATURA, RE_PROCESSO, parece_qualificacao
    from .comum import ROOT, agora_iso, carregar_json, salvar_json
except ImportError:
    from classificar_regras import RE_ASSINATURA, RE_PROCESSO, parece_qualificacao
    from comum import ROOT, agora_iso, carregar_json, salvar_json


def _todos_paragrafos(doc):
    for i, p in enumerate(doc.paragraphs, 1): yield f"corpo:{i}", p
    vistos = []
    def tabela(tabelas, prefixo):
        for ti, t in enumerate(tabelas, 1):
            for ri, r in enumerate(t.rows, 1):
                for ci, c in enumerate(r.cells, 1):
                    if any(c._tc is tc for tc in vistos):
                        continue
                    vistos.append(c._tc)
                    for pi, p in enumerate(c.paragraphs, 1): yield f"{prefixo}:t{ti}l{ri}c{ci}p{pi}", p
                    yield from tabela(c.tables, f"{prefixo}:t{ti}l{ri}c{ci}")
    yield from tabela(doc.tables, "corpo")
    for si, s in enumerate(doc.sections, 1):
        for nome, parte in (("cabecalho", s.header), ("rodape", s.footer)):
            for pi, p in enumerate(parte.paragraphs, 1): yield f"secao{si}:{nome}:p{pi}", p
            yield from tabela(parte.tables, f"secao{si}:{nome}")


def validar(docx: Path, config: dict, aplicacao: dict | None = None) -> dict:
    doc = Document(docx)
    permitidos = set(config["estilos_permitidos"])
    problemas, estilos = [], Counter()
    paragrafos_corpo = doc.paragraphs
    total_corpo = max(1, len(paragrafos_corpo))
    for indice, p in enumerate(paragrafos_corpo):
        texto = " ".join(p.text.split())
        if not texto:
            continue
        estilo = p.style.name if p.style else None
        referencia = f"corpo:{indice + 1}"
        if RE_PROCESSO.search(texto) and len(texto.split()) <= 16 and estilo != "FB - Identificação do Processo":
            problemas.append({
                "severidade": "erro",
                "paragrafo": referencia,
                "codigo": "identificacao_processo_com_estilo_incorreto",
                "detalhe": estilo,
            })
        if indice / total_corpo < 0.30 and parece_qualificacao(texto) and estilo == "FB - Citação":
            problemas.append({
                "severidade": "erro",
                "paragrafo": referencia,
                "codigo": "qualificacao_com_estilo_citacao",
            })
        if (
            indice / total_corpo >= 0.70
            and len(texto.split()) <= 18
            and RE_ASSINATURA.search(texto)
            and estilo != "FB - Assinatura"
        ):
            problemas.append({
                "severidade": "erro",
                "paragrafo": referencia,
                "codigo": "assinatura_com_estilo_incorreto",
                "detalhe": estilo,
            })
        if estilo in {"FB - Título Principal", "FB - Título Secundário"} and indice / total_corpo >= 0.70:
            seguinte = next(
                (" ".join(paragrafos_corpo[j].text.split()) for j in range(indice + 1, len(paragrafos_corpo)) if paragrafos_corpo[j].text.strip()),
                "",
            )
            if 2 <= len(texto.split()) <= 10 and RE_ASSINATURA.search(seguinte):
                problemas.append({
                    "severidade": "erro",
                    "paragrafo": referencia,
                    "codigo": "nome_assinatura_com_estilo_incorreto",
                    "detalhe": estilo,
                })
    vazios = 0
    for i, p in _todos_paragrafos(doc):
        estilo = p.style.name if p.style else None
        estilos[estilo or "SEM ESTILO"] += 1
        texto = p.text.strip()
        if estilo not in permitidos and not str(i).startswith("secao"):
            problemas.append({"severidade": "erro", "paragrafo": i, "codigo": "estilo_nao_autorizado", "detalhe": estilo})
        if not texto:
            if not str(i).startswith("corpo:"):
                vazios = 0
                continue
            vazios += 1
            if vazios >= 2:
                problemas.append({"severidade": "aviso", "paragrafo": i, "codigo": "vazios_repetidos"})
            continue
        vazios = 0
        if "  " in texto:
            problemas.append({"severidade": "aviso", "paragrafo": i, "codigo": "espacos_duplos"})
        if estilo in {"FB - Título Principal", "FB - Título Secundário"} and len(texto.split()) > 18:
            problemas.append({"severidade": "aviso", "paragrafo": i, "codigo": "titulo_longo"})
        if estilo == "FB - Corpo" and str(p.alignment).endswith("CENTER (1)"):
            problemas.append({"severidade": "erro", "paragrafo": i, "codigo": "corpo_centralizado"})
        if estilo == "FB - Citação" and (p.paragraph_format.left_indent and p.paragraph_format.left_indent.cm < 3):
            problemas.append({"severidade": "erro", "paragrafo": i, "codigo": "citacao_sem_recuo"})
        if re.match(r"^\s*\d+[.)]\s", texto) and estilo not in {"FB - Lista", "FB - Pedido", "FB - Título Principal", "FB - Título Secundário"}:
            problemas.append({"severidade": "aviso", "paragrafo": i, "codigo": "lista_sem_estilo_lista"})
        if re.match(r"^\s*(?:###|##|>)\s?", texto) or "**" in texto or re.match(r"^\s*\|.*\|\s*$", texto):
            problemas.append({"severidade": "erro", "paragrafo": i, "codigo": "marcacao_textual_residual"})
    for n, sec in enumerate(doc.sections, 1):
        margens = [sec.top_margin.cm, sec.bottom_margin.cm, sec.left_margin.cm, sec.right_margin.cm]
        if any(x < 1.5 or x > 4.0 for x in margens):
            problemas.append({"severidade": "aviso", "secao": n, "codigo": "margens_suspeitas", "detalhe": margens})
    for ti, tabela in enumerate(doc.tables, 1):
        if len(tabela.rows) == 1 and len(tabela.columns) == 1 and re.search(r"\bOAB(?:/\w+)?\b", tabela.cell(0, 0).text, re.I):
            for pi, p in enumerate(tabela.cell(0, 0).paragraphs, 1):
                if p.style.name != "FB - Assinatura":
                    problemas.append({"severidade": "erro", "paragrafo": f"tabela:{ti}:p{pi}", "codigo": "assinatura_em_tabela_sem_estilo"})
                antes = p.paragraph_format.space_before or p.style.paragraph_format.space_before
                depois = p.paragraph_format.space_after or p.style.paragraph_format.space_after
                if (antes and antes.pt) or (depois and depois.pt):
                    problemas.append({"severidade": "erro", "paragrafo": f"tabela:{ti}:p{pi}", "codigo": "assinatura_com_espacamento"})
    erros = sum(p["severidade"] == "erro" for p in problemas)
    return {
        "documento": str(docx.resolve()), "validado_em": agora_iso(), "total_paragrafos": len(doc.paragraphs),
        "estilos": dict(estilos), "problemas": problemas,
        "resumo": {"erros": erros, "avisos": len(problemas) - erros, "aprovado": erros == 0},
        "aplicacao": aplicacao or {},
    }


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("docx", type=Path)
    ap.add_argument("--config", type=Path, default=ROOT / "modelos" / "estilos.json")
    ap.add_argument("--saida", type=Path, default=ROOT / "trabalho" / "auditoria.json")
    args = ap.parse_args(); salvar_json(args.saida, validar(args.docx, carregar_json(args.config)))


if __name__ == "__main__": main()
