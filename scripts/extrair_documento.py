from __future__ import annotations

import argparse
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn

try:
    from .comum import ROOT, TIPO_PECA_SUPORTADO, agora_iso, hash_texto, salvar_json
except ImportError:
    from comum import ROOT, TIPO_PECA_SUPORTADO, agora_iso, hash_texto, salvar_json


def _pt(valor):
    return round(valor.pt, 2) if valor is not None else None


def _cm(valor):
    return round(valor.cm, 3) if valor is not None else None


def _formato_efetivo(paragrafo, atributo):
    """Usa a formatação direta e, na ausência dela, a definida no estilo."""
    valor = getattr(paragrafo.paragraph_format, atributo)
    if valor is None and paragrafo.style is not None:
        valor = getattr(paragrafo.style.paragraph_format, atributo)
    return valor


def _caracteristicas_runs(paragrafo):
    runs = [r for r in paragrafo.runs if r.text]
    tamanhos = [round(r.font.size.pt, 1) for r in runs if r.font.size]
    return {
        "negrito": bool(runs) and all(bool(r.bold) for r in runs),
        "italico": bool(runs) and all(bool(r.italic) for r in runs),
        "tamanhos_fonte_pt": sorted(set(tamanhos)),
        "formatacao_direta_runs": sum(1 for r in runs if r._element.rPr is not None),
    }


def extrair(caminho: Path) -> dict:
    doc = Document(caminho)
    total = max(1, len(doc.paragraphs))
    pagina = 1
    paragrafos = []
    for i, p in enumerate(doc.paragraphs, 1):
        xml = p._p.xml
        numpr = p._p.pPr.find(qn("w:numPr")) if p._p.pPr is not None else None
        numid = numpr.find(qn("w:numId")) if numpr is not None else None
        ilvl = numpr.find(qn("w:ilvl")) if numpr is not None else None
        registro = {
            "id": i,
            "texto": p.text,
            "hash": hash_texto(p.text),
            "estilo_original": p.style.name if p.style else None,
            "alinhamento": str(p.alignment) if p.alignment is not None else None,
            "recuo_esquerdo_cm": _cm(_formato_efetivo(p, "left_indent")),
            "recuo_primeira_linha_cm": _cm(_formato_efetivo(p, "first_line_indent")),
            "espacamento_antes_pt": _pt(_formato_efetivo(p, "space_before")),
            "espacamento_depois_pt": _pt(_formato_efetivo(p, "space_after")),
            "possui_numeracao": numpr is not None and numid is not None,
            "numero_lista_id": numid.get(qn("w:val")) if numid is not None else None,
            "nivel_lista": int(ilvl.get(qn("w:val"))) if ilvl is not None else 0,
            "posicao_relativa": round((i - 1) / total, 4),
            "pagina_aproximada": pagina,
            **_caracteristicas_runs(p),
        }
        paragrafos.append(registro)
        pagina += xml.count('w:type="page"') + xml.count("w:lastRenderedPageBreak")

    tabelas = []
    for i, tabela in enumerate(doc.tables, 1):
        tabelas.append({"id": i, "linhas": len(tabela.rows), "colunas": len(tabela.columns), "texto": [[c.text for c in r.cells] for r in tabela.rows]})
    cabecalhos, rodapes = [], []
    for i, sec in enumerate(doc.sections, 1):
        cabecalhos.append({"secao": i, "paragrafos": [p.text for p in sec.header.paragraphs]})
        rodapes.append({"secao": i, "paragrafos": [p.text for p in sec.footer.paragraphs]})
    with zipfile.ZipFile(caminho) as zf:
        nomes = set(zf.namelist())
        imagens = [n for n in nomes if n.startswith("word/media/")]
        notas = "word/footnotes.xml" in nomes
        campos = sum(zf.read(n).count(b"w:fldChar") for n in nomes if n.startswith("word/") and n.endswith(".xml"))
    return {
        "documento": caminho.name,
        "caminho": str(caminho.resolve()),
        "extraido_em": agora_iso(),
        "tipo_peca": TIPO_PECA_SUPORTADO,
        "paragrafos": paragrafos,
        "tabelas": tabelas,
        "cabecalhos": cabecalhos,
        "rodapes": rodapes,
        "secoes": len(doc.sections),
        "imagens": imagens,
        "possui_notas_rodape": notas,
        "marcadores_de_campos": campos,
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("docx", type=Path)
    ap.add_argument("--saida", type=Path, default=ROOT / "trabalho" / "estrutura_documento.json")
    args = ap.parse_args()
    salvar_json(args.saida, extrair(args.docx))
    print(args.saida)


if __name__ == "__main__":
    main()
