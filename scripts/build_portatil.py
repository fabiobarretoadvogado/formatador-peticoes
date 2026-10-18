from __future__ import annotations

import importlib.util
import json
import os
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app_version import VERSION
from scripts.gerar_icone import main as gerar_icone
from scripts.pacote_atualizacao import gerar_pacote, manifesto_publicado
from scripts.protocolo_atualizacao import MANIFESTO, sha256


ROOT = Path(__file__).resolve().parents[1]
DIST = ROOT / "dist"


def main() -> None:
    gerar_icone()
    if str(ROOT) not in sys.path:
        sys.path.insert(0, str(ROOT))
    dependencias_ausentes = [
        modulo
        for modulo in ("PyInstaller", "markdown_it", "scripts.numeracao_word")
        if importlib.util.find_spec(modulo) is None
    ]
    if dependencias_ausentes:
        raise SystemExit(
            "Dependências do pacote portátil não estão disponíveis neste Python: "
            + ", ".join(dependencias_ausentes)
            + ". Execute:\n"
            f'  "{sys.executable}" -m pip install -e ".[build]"'
        )

    with tempfile.TemporaryDirectory(prefix="formatador_portatil_") as temporario:
        temporario_raiz = Path(temporario)
        dist_temporario = temporario_raiz / "dist"
        build_temporario = temporario_raiz / "build"
        pacote = dist_temporario / "FormatadorPeticoes"

        subprocess.run(
            [
                sys.executable,
                "-m",
                "PyInstaller",
                "--noconfirm",
                "--clean",
                "--distpath",
                str(dist_temporario),
                "--workpath",
                str(build_temporario),
                str(ROOT / "FormatadorPeticoes.spec"),
            ],
            cwd=ROOT,
            check=True,
        )

        # O iniciador permanece fixo; cada versão completa ocupa uma pasta.
        aplicativo = dist_temporario / "aplicativo_compilado"
        pacote.rename(aplicativo)
        versoes = pacote / "versoes"
        versoes.mkdir(parents=True)
        aplicativo = aplicativo.rename(versoes / VERSION)
        modelos_destino = aplicativo / "modelos"
        modelos_destino.mkdir(parents=True, exist_ok=True)
        for nome in ("modelo_oficial.docx", "estilos.json"):
            shutil.copy2(ROOT / "modelos" / nome, modelos_destino / nome)

        gerar_pacote(aplicativo, DIST, VERSION, manifesto_publicado())
        (pacote / "versao_inicial.json").write_text(json.dumps({"versao": VERSION}), encoding="utf-8")
        subprocess.run([
            sys.executable, "-m", "PyInstaller", "--noconfirm", "--clean", "--onefile", "--windowed",
            "--name", "FormatadorPeticoes", "--paths", str(ROOT),
            "--distpath", str(pacote), "--workpath", str(temporario_raiz / "build_iniciador"),
            "--specpath", str(temporario_raiz), "--icon", str(ROOT / "assets/aplicativo.ico"),
            str(ROOT / "scripts/iniciar_aplicativo.py"),
        ], cwd=ROOT, check=True)

        saida_teste = temporario_raiz / "teste_resultados"
        entrada_markdown = temporario_raiz / "peticao_markdown.md"
        entrada_markdown.write_text(
            "## DOS FATOS\n\nTexto de teste do pacote portátil.\n\n"
            "> Citação de teste.\n\n1. Item numerado.\n\n"
            "| Campo | Valor |\n| --- | --- |\n| Origem | Markdown |\n",
            encoding="utf-8",
        )
        entrada_word = temporario_raiz / "peticao_word_sem_timbre.docx"
        controle_word = Document()
        controle_word.sections[0].top_margin = Cm(3)
        controle_word.sections[0].bottom_margin = Cm(2)
        controle_word.sections[0].left_margin = Cm(3)
        controle_word.sections[0].right_margin = Cm(2)
        controle_word.add_paragraph("Texto de teste para conferir o papel timbrado do modelo.")
        controle_word.add_paragraph("Título principal de controle", style="Heading 1").runs[0].bold = False
        subtitulo = controle_word.add_paragraph("Subtítulo de controle", style="Heading 2")
        subtitulo.runs[0].bold = False
        subtitulo.paragraph_format.first_line_indent = Cm(-1)
        for texto in (
            "Excelentíssimo Senhor Juiz de Direito da Vara Cível",
            "Processo nº 0001234-56.2025.8.26.0001",
        ):
            run = controle_word.add_paragraph(texto).runs[0]
            run.bold = run.font.all_caps = False
        controle_word.sections[0].footer.paragraphs[0].text = "Rodapé original de controle"
        controle_word.save(entrada_word)
        ambiente = os.environ.copy()
        ambiente["LOCALAPPDATA"] = str(temporario_raiz / "dados_locais")
        try:
            teste = subprocess.run(
                [
                    str(pacote / "FormatadorPeticoes.exe"),
                    str(ROOT / "modelos" / "modelo_oficial.docx"),
                    str(entrada_markdown),
                    str(entrada_word),
                    "--saida",
                    str(saida_teste),
                    "--sem-pdf",
                    "--silencioso",
                ],
                cwd=pacote,
                env=ambiente,
                check=False,
                timeout=90,
            )
        except subprocess.TimeoutExpired as exc:
            raise RuntimeError(
                "O executável foi criado, mas o teste de DOCX e Markdown não terminou em 90 segundos."
            ) from exc
        saidas_docx = list(saida_teste.rglob("*_formatada.docx"))
        if (
            teste.returncode != 0
            or len(saidas_docx) < 3
            or not any(saida_teste.rglob("entrada_intermediaria.docx"))
            or not any(saida_teste.rglob("manifesto_entrada.json"))
        ):
            raise RuntimeError(
                "O executável foi criado, mas não passou no teste automático de DOCX e Markdown."
            )

        saidas_word = list(saida_teste.rglob("peticao_word_sem_timbre_formatada.docx"))
        if len(saidas_word) != 1:
            raise RuntimeError("O executável não produziu o Word de controle do papel timbrado.")
        word_final = Document(saidas_word[0])
        cabecalho = word_final.sections[0].header
        imagens_timbre = cabecalho._element.xpath(".//a:blip")
        if not imagens_timbre:
            raise RuntimeError("O executável não acrescentou o timbre do modelo à entrada Word.")
        for imagem in imagens_timbre:
            rel = cabecalho.part.rels[imagem.get(qn("r:embed"))]
            if not rel.target_part.blob:
                raise RuntimeError("A imagem timbrada do Word está vazia.")
        if word_final.sections[0].footer.paragraphs[0].text != "Rodapé original de controle":
            raise RuntimeError("O executável alterou o rodapé original da entrada Word.")
        if word_final.paragraphs[0].text != controle_word.paragraphs[0].text:
            raise RuntimeError("O executável alterou o texto do Word de controle do timbre.")
        referencia = Document(ROOT / "modelos" / "modelo_oficial.docx").sections[0]
        for secao in word_final.sections:
            if secao._sectPr.pgMar.attrib != referencia._sectPr.pgMar.attrib:
                raise RuntimeError("O executável não aplicou as margens do modelo oficial à entrada Word.")
        print("Timbre do modelo aplicado ao Word sem cabeçalho; texto e rodapé preservados.")
        print("Margens e distâncias de cabeçalho/rodapé iguais às do modelo oficial.")
        for indice, estilo in ((1, "FB - Título Principal"), (2, "FB - Título Secundário")):
            titulo = word_final.paragraphs[indice]
            if titulo.style.name != estilo or not all(run.bold for run in titulo.runs):
                raise RuntimeError("O executável não garantiu o negrito dos títulos.")
        subtitulo_final = word_final.paragraphs[2]
        if abs(subtitulo_final.paragraph_format.first_line_indent.cm - 1) > 0.01:
            raise RuntimeError("O executável não aplicou primeira linha de 1 cm ao subtítulo.")
        print("Títulos em negrito e subtítulo com primeira linha de 1 cm validados.")
        for indice, estilo in ((3, "FB - Endereçamento"), (4, "FB - Identificação do Processo")):
            paragrafo = word_final.paragraphs[indice]
            if paragrafo.style.name != estilo or not all(
                run.bold and run.font.all_caps for run in paragrafo.runs
            ):
                raise RuntimeError("O executável não garantiu negrito e caixa alta no endereçamento/processo.")
            if paragrafo.text != controle_word.paragraphs[indice].text:
                raise RuntimeError("O executável alterou o texto original do endereçamento/processo.")
        print("Endereçamento e processo em negrito e caixa alta visual validados.")

        diagnostico_interface = temporario_raiz / "diagnostico_interface.json"
        try:
            teste_interface = subprocess.run(
                [
                    str(pacote / "FormatadorPeticoes.exe"),
                    "--validar-interface",
                    str(diagnostico_interface),
                ],
                cwd=pacote,
                env=ambiente,
                check=False,
                timeout=30,
            )
        except subprocess.TimeoutExpired as exc:
            raise RuntimeError(
                "O executável processa documentos, mas a validação da interface não terminou em 30 segundos."
            ) from exc
        if teste_interface.returncode != 0 or not diagnostico_interface.is_file():
            detalhe = (
                diagnostico_interface.read_text(encoding="utf-8")
                if diagnostico_interface.is_file()
                else "O executável não produziu o diagnóstico esperado."
            )
            raise RuntimeError(
                "O executável processa documentos, mas o frontend gráfico não iniciou:\n" + detalhe
            )

        diagnostico = json.loads(diagnostico_interface.read_text(encoding="utf-8"))
        if not diagnostico.get("tk_importado") or not diagnostico.get("funcional"):
            raise RuntimeError(
                "O pacote seria publicado sem o frontend completo:\n"
                + diagnostico_interface.read_text(encoding="utf-8")
            )
        print("Frontend Tcl/Tk completo validado.")

        # Exercita o iniciador compilado e a confirmação real pelo frontend.
        teste_update = temporario_raiz / "teste_atualizacao"
        teste_update.mkdir()
        shutil.copy2(pacote / "FormatadorPeticoes.exe", teste_update / "FormatadorPeticoes.exe")
        shutil.copytree(aplicativo, teste_update / "versoes" / VERSION)
        anterior = teste_update / "versoes" / "0.0.0"
        shutil.copytree(aplicativo, anterior)
        dados_anteriores = json.loads((anterior / MANIFESTO).read_text(encoding="utf-8"))
        dados_anteriores["versao"] = "0.0.0"
        for item in dados_anteriores["arquivos"]:
            item["pacote"] = "0.0.0"
        (anterior / MANIFESTO).write_text(json.dumps(dados_anteriores), encoding="utf-8")
        (teste_update / "versao_inicial.json").write_text(json.dumps({"versao": "0.0.0"}), encoding="utf-8")
        diagnostico_update = temporario_raiz / "diagnostico_atualizacao.json"
        subprocess.run([
            str(teste_update / "FormatadorPeticoes.exe"), "--ativar", VERSION,
            "--manifesto-sha256", sha256(teste_update / "versoes" / VERSION / MANIFESTO),
            "--validar-ativacao", str(diagnostico_update),
        ], env=ambiente, cwd=teste_update, timeout=180, check=True)
        estado = json.loads((teste_update / "versao_ativa.json").read_text(encoding="utf-8"))
        if estado.get("versao") != VERSION or not diagnostico_update.is_file():
            raise RuntimeError("O iniciador não confirmou a troca para a nova versão.")
        print("Atualização e reinício pelo iniciador compilado validados.")

        zip_base = temporario_raiz / "FormatadorPeticoes-portatil"
        zip_criado = Path(
            shutil.make_archive(
                str(zip_base), "zip", root_dir=dist_temporario, base_dir=pacote.name
            )
        )

        DIST.mkdir(parents=True, exist_ok=True)
        destino = DIST / f"FormatadorPeticoes-portatil-{VERSION}.zip"
        shutil.copy2(zip_criado, destino)
        print(f"Pacote criado e testado: {destino}")


if __name__ == "__main__":
    main()
