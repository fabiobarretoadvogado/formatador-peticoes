from __future__ import annotations

import hashlib
import json
import re
import shutil
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import unquote, urlsplit

import markdown_it
from docx import Document
from docx.document import Document as DocumentoDocx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from markdown_it import MarkdownIt
from markdown_it.tree import SyntaxTreeNode
from PIL import Image

try:
    from .comum import agora_iso, salvar_json
except ImportError:
    from comum import agora_iso, salvar_json


EXTENSOES_IMAGEM = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff"}
MARCADOR_LETRA = "\ue000MDLETTER:{caixa}:{numero}\ue001"
RE_MARCADOR_LETRA = re.compile(r"\ue000MDLETTER:([LU]):(\d+)\ue001")
RE_LISTA_LETRA = re.compile(r"^([A-Za-z])[.)]\s+(.+)$")
RE_NOTA_RODAPE = re.compile(r"\[\^[^\]]+\]")
PERFIL_MARKDOWN_PADRAO = {
    "versao": 1,
    "codificacoes": ["utf-8", "utf-8-sig"],
    "titulos": ["##", "###"],
    "citacoes_niveis": 1,
    "listas_niveis": 1,
    "listas": ["numerica", "alfabetica", "marcadores"],
    "tabelas": "gfm_simples",
    "negrito": "intermediario_preserva_saida_segue_estilos",
    "italico": True,
    "imagens": "locais_relativas_isoladas_em_bloco",
    "html": False,
    "links": False,
    "notas_rodape": False,
    "blocos_codigo": False,
}


class RecursoMarkdownNaoSuportado(ValueError):
    pass


@dataclass(frozen=True)
class ResultadoIngestao:
    origem: Path
    formato_original: str
    entrada_docx: Path
    docx_intermediario: Path | None
    manifesto: Path
    dados: dict


def _sha256(caminho: Path) -> str:
    digest = hashlib.sha256()
    with caminho.open("rb") as stream:
        for bloco in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(bloco)
    return digest.hexdigest()


def _hash_estrutura(estrutura: list[dict]) -> str:
    serializado = json.dumps(estrutura, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(serializado.encode("utf-8")).hexdigest()


def _limpar_corpo(doc: DocumentoDocx) -> None:
    corpo = doc._body._element
    for elemento in list(corpo):
        if not elemento.tag.endswith("}sectPr"):
            corpo.remove(elemento)


def _filho(nome: str, valor: str | None = None):
    no = OxmlElement(f"w:{nome}")
    if valor is not None:
        no.set(qn("w:val"), valor)
    return no


def _proximo_id(elemento, nome: str, atributo: str) -> str:
    usados = [
        int(no.get(qn(atributo)))
        for no in elemento.findall(qn(nome))
        if no.get(qn(atributo), "").isdigit()
    ]
    return str(max(usados, default=0) + 1)


def _criar_numeracao(doc: DocumentoDocx, formato: str, inicio: int = 1) -> str:
    numeracao = doc.part.numbering_part.element
    abstract_id = _proximo_id(numeracao, "w:abstractNum", "w:abstractNumId")
    num_id = _proximo_id(numeracao, "w:num", "w:numId")

    abstrato = OxmlElement("w:abstractNum")
    abstrato.set(qn("w:abstractNumId"), abstract_id)
    abstrato.append(_filho("multiLevelType", "singleLevel"))
    nivel = OxmlElement("w:lvl")
    nivel.set(qn("w:ilvl"), "0")
    nivel.append(_filho("start", str(inicio)))
    nivel.append(_filho("numFmt", formato))
    nivel.append(_filho("lvlText", "•" if formato == "bullet" else "%1."))
    nivel.append(_filho("lvlJc", "left"))
    ppr = OxmlElement("w:pPr")
    recuo = OxmlElement("w:ind")
    recuo.set(qn("w:left"), "1701")
    recuo.set(qn("w:hanging"), "425")
    ppr.append(recuo)
    nivel.append(ppr)
    abstrato.append(nivel)
    numeracao.append(abstrato)

    lista = OxmlElement("w:num")
    lista.set(qn("w:numId"), num_id)
    lista.append(_filho("abstractNumId", abstract_id))
    numeracao.append(lista)
    return num_id


def _aplicar_numeracao(paragrafo: Paragraph, num_id: str) -> None:
    ppr = paragrafo._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    numpr.append(_filho("ilvl", "0"))
    numpr.append(_filho("numId", num_id))
    ppr.append(numpr)


def _preprocessar_listas_alfabeticas(texto: str) -> tuple[str, int]:
    linhas = []
    convertidas = 0
    for linha in texto.splitlines():
        encontrado = RE_LISTA_LETRA.match(linha)
        if not encontrado:
            linhas.append(linha)
            continue
        letra, conteudo = encontrado.groups()
        numero = ord(letra.casefold()) - ord("a") + 1
        if not 1 <= numero <= 26:
            linhas.append(linha)
            continue
        caixa = "U" if letra.isupper() else "L"
        marcador = MARCADOR_LETRA.format(caixa=caixa, numero=numero)
        linhas.append(f"{numero}. {marcador}{conteudo}")
        convertidas += 1
    return "\n".join(linhas), convertidas


def _linha_no(node: SyntaxTreeNode) -> int | None:
    if node.map:
        return int(node.map[0]) + 1
    return None


def _erro_recurso(node: SyntaxTreeNode, recurso: str) -> RecursoMarkdownNaoSuportado:
    linha = _linha_no(node)
    local = f" na linha {linha}" if linha else ""
    return RecursoMarkdownNaoSuportado(f"Recurso Markdown não suportado{local}: {recurso}")


def _texto_node_bruto(node: SyntaxTreeNode) -> str:
    if node.type == "text":
        return node.content
    return "".join(_texto_node_bruto(filho) for filho in node.children)


def _texto_node(node: SyntaxTreeNode) -> str:
    return RE_MARCADOR_LETRA.sub("", _texto_node_bruto(node))


def _adicionar_texto(paragrafo: Paragraph, texto: str, *, negrito: bool, italico: bool) -> None:
    texto = RE_MARCADOR_LETRA.sub("", texto)
    if not texto:
        return
    run = paragrafo.add_run(texto)
    if negrito:
        run.bold = True
    if italico:
        run.italic = True


def _renderizar_inline(
    node: SyntaxTreeNode,
    paragrafo: Paragraph,
    *,
    negrito: bool = False,
    italico: bool = False,
) -> None:
    if node.type == "text":
        _adicionar_texto(paragrafo, node.content, negrito=negrito, italico=italico)
        return
    if node.type == "softbreak":
        _adicionar_texto(paragrafo, " ", negrito=negrito, italico=italico)
        return
    if node.type == "hardbreak":
        paragrafo.add_run().add_break()
        return
    if node.type == "strong":
        for filho in node.children:
            _renderizar_inline(filho, paragrafo, negrito=True, italico=italico)
        return
    if node.type == "em":
        for filho in node.children:
            _renderizar_inline(filho, paragrafo, negrito=negrito, italico=True)
        return
    if node.type == "inline":
        for filho in node.children:
            _renderizar_inline(filho, paragrafo, negrito=negrito, italico=italico)
        return
    if node.type == "image":
        raise _erro_recurso(node, "imagem misturada com texto no mesmo parágrafo")
    if node.type == "link":
        raise _erro_recurso(node, "links; use texto simples na primeira versão")
    if node.type in {"code_inline", "html_inline", "s"}:
        raise _erro_recurso(node, node.type)
    raise _erro_recurso(node, node.type)


def _inline_do_bloco(node: SyntaxTreeNode) -> SyntaxTreeNode:
    if len(node.children) != 1 or node.children[0].type != "inline":
        raise _erro_recurso(node, "estrutura interna inesperada")
    return node.children[0]


def _imagem_isolada(node: SyntaxTreeNode) -> SyntaxTreeNode | None:
    inline = _inline_do_bloco(node)
    significativos = [
        filho
        for filho in inline.children
        if filho.type not in {"softbreak", "hardbreak"}
        and not (filho.type == "text" and not filho.content.strip())
    ]
    if len(significativos) == 1 and significativos[0].type == "image":
        return significativos[0]
    if any(filho.type == "image" for filho in significativos):
        raise _erro_recurso(node, "imagem misturada com texto no mesmo parágrafo")
    return None


def _resolver_imagem(referencia: str, pasta_markdown: Path) -> Path:
    partes = urlsplit(referencia)
    if partes.scheme or partes.netloc:
        raise RecursoMarkdownNaoSuportado(
            f"Somente imagens locais relativas ao Markdown são aceitas: {referencia}"
        )
    # URLs Markdown usam barras normais; converta-as sem transformar a
    # referência em caminho absoluto.
    relativo = Path(unquote(partes.path).replace("/", "\\"))
    if relativo.is_absolute() or ".." in relativo.parts:
        raise RecursoMarkdownNaoSuportado(
            f"A imagem deve permanecer dentro da pasta do Markdown: {referencia}"
        )
    base = pasta_markdown.resolve()
    caminho = (base / relativo).resolve()
    if not caminho.is_relative_to(base):
        raise RecursoMarkdownNaoSuportado(
            f"A imagem deve permanecer dentro da pasta do Markdown: {referencia}"
        )
    if caminho.suffix.casefold() not in EXTENSOES_IMAGEM:
        raise RecursoMarkdownNaoSuportado(
            f"Formato de imagem não suportado: {caminho.suffix or referencia}"
        )
    if not caminho.is_file():
        raise FileNotFoundError(f"Imagem referenciada no Markdown não encontrada: {caminho}")
    try:
        with Image.open(caminho) as imagem:
            imagem.verify()
    except Exception as exc:
        raise ValueError(f"A imagem referenciada não pôde ser validada: {caminho}") from exc
    return caminho


def _adicionar_imagem(
    doc: DocumentoDocx,
    node: SyntaxTreeNode,
    pasta_markdown: Path,
    imagens: list[dict],
) -> Paragraph:
    referencia = str(node.attrs.get("src") or "")
    caminho = _resolver_imagem(referencia, pasta_markdown)
    descricao = _texto_node(node).strip()
    paragrafo = doc.add_paragraph(style="Normal")
    paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = paragrafo.add_run().add_picture(str(caminho))
    secao = doc.sections[-1]
    largura_disponivel = secao.page_width - secao.left_margin - secao.right_margin
    if shape.width > largura_disponivel:
        proporcao = largura_disponivel / shape.width
        shape.width = largura_disponivel
        shape.height = int(shape.height * proporcao)
    if descricao:
        shape._inline.docPr.set("descr", descricao)
    imagens.append({
        "referencia": referencia,
        "arquivo": str(caminho),
        "sha256": _sha256(caminho),
        "descricao": descricao,
        "largura_emu": int(shape.width),
        "altura_emu": int(shape.height),
    })
    return paragrafo


def _renderizar_paragrafo(doc: DocumentoDocx, node: SyntaxTreeNode, estilo: str = "Normal") -> Paragraph:
    paragrafo = doc.add_paragraph(style=estilo)
    _renderizar_inline(_inline_do_bloco(node), paragrafo)
    return paragrafo


def _renderizar_lista(doc: DocumentoDocx, node: SyntaxTreeNode, contagem: dict) -> None:
    itens = node.children
    if not itens:
        return
    for item in itens:
        if item.type != "list_item" or len(item.children) != 1 or item.children[0].type != "paragraph":
            raise _erro_recurso(node, "listas aninhadas ou itens com vários parágrafos")
    formato_atual = None
    num_id = None
    inicio_lista = int(node.attrs.get("start", 1))
    for indice, item in enumerate(itens):
        marcador = RE_MARCADOR_LETRA.search(_texto_node_bruto(item))
        if node.type == "bullet_list":
            formato, inicio, chave_contagem = "bullet", 1, "listas_marcadores"
        elif marcador:
            formato = "upperLetter" if marcador.group(1) == "U" else "lowerLetter"
            inicio, chave_contagem = int(marcador.group(2)), "listas_alfabeticas"
        else:
            formato, inicio, chave_contagem = "decimal", inicio_lista + indice, "listas_numericas"
        if formato != formato_atual:
            num_id = _criar_numeracao(doc, formato, inicio)
            contagem[chave_contagem] += 1
            formato_atual = formato
        paragrafo = _renderizar_paragrafo(doc, item.children[0])
        assert num_id is not None
        _aplicar_numeracao(paragrafo, num_id)
        contagem["itens_lista"] += 1


def _linhas_tabela(node: SyntaxTreeNode) -> list[tuple[bool, list[SyntaxTreeNode]]]:
    linhas: list[tuple[bool, list[SyntaxTreeNode]]] = []
    for secao in node.children:
        if secao.type not in {"thead", "tbody"}:
            raise _erro_recurso(secao, secao.type)
        cabecalho = secao.type == "thead"
        for linha in secao.children:
            if linha.type != "tr":
                raise _erro_recurso(linha, linha.type)
            linhas.append((cabecalho, list(linha.children)))
    return linhas


def _renderizar_tabela(doc: DocumentoDocx, node: SyntaxTreeNode, contagem: dict) -> Table:
    linhas = _linhas_tabela(node)
    colunas = max((len(celulas) for _, celulas in linhas), default=0)
    if not linhas or not colunas:
        raise _erro_recurso(node, "tabela vazia")
    if any(len(celulas) != colunas for _, celulas in linhas):
        raise _erro_recurso(node, "tabela com quantidade irregular de células")
    tabela = doc.add_table(rows=len(linhas), cols=colunas)
    if "Table Grid" in [estilo.name for estilo in doc.styles]:
        tabela.style = "Table Grid"
    for ri, (cabecalho, celulas) in enumerate(linhas):
        for ci, celula_node in enumerate(celulas):
            if celula_node.type not in {"th", "td"}:
                raise _erro_recurso(celula_node, celula_node.type)
            paragrafo = tabela.cell(ri, ci).paragraphs[0]
            inline = _inline_do_bloco(celula_node)
            _renderizar_inline(inline, paragrafo, negrito=cabecalho)
    contagem["tabelas"] += 1
    contagem["celulas_tabela"] += len(linhas) * colunas
    return tabela


def _estrutura_docx(doc: DocumentoDocx) -> list[dict]:
    estrutura: list[dict] = []
    for elemento in doc.element.body.iterchildren():
        if isinstance(elemento, CT_P):
            paragrafo = Paragraph(elemento, doc)
            desenhos = paragrafo._p.xpath(".//w:drawing")
            if desenhos:
                descricoes = [
                    no.get("descr", "")
                    for no in paragrafo._p.xpath(".//wp:docPr")
                ]
                estrutura.append({"tipo": "imagem", "descricoes": descricoes})
            else:
                estrutura.append({"tipo": "paragrafo", "texto": paragrafo.text})
        elif isinstance(elemento, CT_Tbl):
            tabela = Table(elemento, doc)
            estrutura.append({
                "tipo": "tabela",
                "celulas": [[celula.text for celula in linha.cells] for linha in tabela.rows],
            })
    return estrutura


def converter_markdown(
    markdown: Path,
    modelo: Path,
    saida: Path,
    manifesto: Path,
    *,
    perfil_markdown: dict | None = None,
) -> dict:
    markdown = markdown.expanduser().resolve()
    modelo = modelo.expanduser().resolve()
    saida = saida.expanduser().resolve()
    manifesto = manifesto.expanduser().resolve()
    texto_original = markdown.read_text(encoding="utf-8-sig")
    if RE_NOTA_RODAPE.search(texto_original):
        raise RecursoMarkdownNaoSuportado("Notas de rodapé Markdown ainda não são suportadas")
    texto, listas_alfabeticas = _preprocessar_listas_alfabeticas(texto_original)
    parser = MarkdownIt("commonmark", {"html": False}).enable("table")
    raiz = SyntaxTreeNode(parser.parse(texto))

    saida.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(modelo, saida)
    doc = Document(saida)
    _limpar_corpo(doc)
    contagem = {
        "titulos_principais": 0,
        "titulos_secundarios": 0,
        "paragrafos": 0,
        "citacoes": 0,
        "listas_numericas": 0,
        "listas_alfabeticas": 0,
        "listas_marcadores": 0,
        "itens_lista": 0,
        "tabelas": 0,
        "celulas_tabela": 0,
        "imagens": 0,
    }
    imagens: list[dict] = []
    for node in raiz.children:
        if node.type == "heading":
            if node.tag == "h2":
                _renderizar_paragrafo(doc, node, "Heading 1")
                contagem["titulos_principais"] += 1
            elif node.tag == "h3":
                _renderizar_paragrafo(doc, node, "Heading 2")
                contagem["titulos_secundarios"] += 1
            else:
                raise _erro_recurso(node, f"nível de título {node.tag}; use apenas ## e ###")
        elif node.type == "paragraph":
            imagem = _imagem_isolada(node)
            if imagem:
                _adicionar_imagem(doc, imagem, markdown.parent, imagens)
                contagem["imagens"] += 1
            else:
                _renderizar_paragrafo(doc, node)
                contagem["paragrafos"] += 1
        elif node.type == "blockquote":
            if not node.children or any(filho.type != "paragraph" for filho in node.children):
                raise _erro_recurso(node, "citações aninhadas ou com blocos complexos")
            for filho in node.children:
                _renderizar_paragrafo(doc, filho, "Quote")
                contagem["citacoes"] += 1
        elif node.type in {"ordered_list", "bullet_list"}:
            _renderizar_lista(doc, node, contagem)
        elif node.type == "table":
            _renderizar_tabela(doc, node, contagem)
        elif node.type in {"html_block", "fence", "code_block", "hr"}:
            raise _erro_recurso(node, node.type)
        else:
            raise _erro_recurso(node, node.type)

    estrutura_esperada = _estrutura_docx(doc)
    doc.save(saida)
    estrutura_gravada = _estrutura_docx(Document(saida))
    hash_esperado = _hash_estrutura(estrutura_esperada)
    hash_gravado = _hash_estrutura(estrutura_gravada)
    if hash_esperado != hash_gravado:
        raise RuntimeError("A auditoria estrutural do DOCX intermediário detectou perda de conteúdo")

    dados = {
        "versao_manifesto": 1,
        "criado_em": agora_iso(),
        "formato_original": "markdown",
        "origem": str(markdown),
        "sha256_origem": _sha256(markdown),
        "entrada_pipeline": str(saida),
        "docx_intermediario": str(saida),
        "sha256_docx_intermediario": _sha256(saida),
        "parser": {
            "nome": "markdown-it-py",
            "versao": markdown_it.__version__,
            "perfil": "commonmark+table",
            "html": False,
        },
        "perfil_markdown": perfil_markdown or PERFIL_MARKDOWN_PADRAO,
        "listas_alfabeticas_preprocessadas": listas_alfabeticas,
        "contagem": contagem,
        "imagens": imagens,
        "avisos": [],
        "auditoria_conteudo": {
            "aprovada": True,
            "hash_estrutura_antes_de_gravar": hash_esperado,
            "hash_estrutura_apos_gravar": hash_gravado,
        },
    }
    salvar_json(manifesto, dados)
    return dados


def normalizar_entrada(
    entrada: Path,
    modelo: Path,
    root: Path,
    *,
    perfil_markdown: dict | None = None,
) -> ResultadoIngestao:
    entrada = entrada.expanduser().resolve()
    manifesto = root / "trabalho" / "manifesto_entrada.json"
    if entrada.suffix.casefold() == ".md":
        intermediario = root / "trabalho" / "entrada_intermediaria.docx"
        dados = converter_markdown(
            entrada,
            modelo,
            intermediario,
            manifesto,
            perfil_markdown=perfil_markdown,
        )
        return ResultadoIngestao(entrada, "markdown", intermediario, intermediario, manifesto, dados)
    if entrada.suffix.casefold() == ".docx":
        dados = {
            "versao_manifesto": 1,
            "criado_em": agora_iso(),
            "formato_original": "docx",
            "origem": str(entrada),
            "sha256_origem": _sha256(entrada),
            "entrada_pipeline": str(entrada),
            "docx_intermediario": None,
            "sha256_docx_intermediario": None,
            "parser": None,
            "contagem": {},
            "imagens": [],
            "avisos": [],
            "auditoria_conteudo": {"aprovada": True, "modo": "entrada DOCX nativa"},
        }
        salvar_json(manifesto, dados)
        return ResultadoIngestao(entrada, "docx", entrada, None, manifesto, dados)
    raise ValueError("Informe uma petição inicial existente em DOCX ou Markdown (.md)")
