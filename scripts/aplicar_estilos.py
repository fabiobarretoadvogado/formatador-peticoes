from __future__ import annotations

import argparse
import copy
import io
import posixpath
import re
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.run import Run
from lxml import etree
from PIL import Image, ImageOps

try:
    from .comum import ROOT, carregar_json
    from .numeracao_word import ESTILOS_TITULO, criar_lista_multinivel_titulos, criar_numpr
except ImportError:
    from comum import ROOT, carregar_json
    from numeracao_word import ESTILOS_TITULO, criar_lista_multinivel_titulos, criar_numpr

PPR_LIMPAR = {"jc", "ind", "spacing", "contextualSpacing", "mirrorIndents", "adjustRightInd", "snapToGrid", "keepLines", "keepNext", "widowControl", "pBdr"}
RPR_LIMPAR = {"rFonts", "sz", "szCs", "color", "spacing", "position", "kern", "caps", "smallCaps", "outline", "shadow", "emboss", "imprint", "fitText", "w", "snapToGrid"}

RE_NOME_ENTIDADE_CAIXA_ALTA = re.compile(
    r"\b(?:[A-ZÁÀÂÃÉÊÍÓÔÕÚÇÜ]{2,}|DA|DAS|DE|DO|DOS|E|EM|NA|NO|A|O)"
    r"(?:\s+(?:[A-ZÁÀÂÃÉÊÍÓÔÕÚÇÜ]{2,}|DA|DAS|DE|DO|DOS|E|EM|NA|NO|A|O)){1,}"
    r"(?=\s|[,.;:]|$)"
)
RE_ACAO_FASE_PROCESSUAL = re.compile(
    r"\b(?:ação(?:s)?|embargos\s+de\s+declaração|sentença|réplica|contestação|"
    r"petição\s+inicial|apelação|agravo(?:\s+de\s+instrumento|\s+interno)?|"
    r"mandado\s+de\s+segurança|cumprimento\s+de\s+sentença|execução|"
    r"procedimento|recurso|decisão|acórdão|despacho)\b",
    re.IGNORECASE,
)
RE_MARKDOWN_ESTRUTURAL = re.compile(r"^\s*(?:####|###|##|>)\s?")
RE_MARKDOWN_TABELA = re.compile(r"^\s*\|.*\|\s*$")
RE_SEPARADOR_TABELA = re.compile(r"^:?-{3,}:?$")
RE_TITULO_NUMERADO_LITERAL = re.compile(
    r"^\s*(?:(?:\d+(?:\.\d+)*|[IVXLCDM]+)[.)]?)(?:\s*[-–—:]\s*|\s+)(\S.*)$",
    re.IGNORECASE,
)
RE_PEDIDO_NUMERADO_LITERAL = re.compile(r"^\s*(\d+)[.)]\s+(.+)$")


def _substituir_texto(paragrafo, texto: str) -> None:
    for run in list(paragrafo._p.findall(qn("w:r"))):
        paragrafo._p.remove(run)
    paragrafo.add_run(texto)


def _normalizar_markdown_paragrafo(paragrafo, estilo: str) -> bool:
    texto = paragrafo.text
    novo = texto
    if estilo in {"FB - Título Principal", "FB - Título Secundário", "FB - Citação"}:
        novo = RE_MARKDOWN_ESTRUTURAL.sub("", novo, count=1)
    novo = novo.replace("**", "")
    if novo != texto:
        _substituir_texto(paragrafo, novo)
        return True
    return False


def _normalizar_titulo_literal(paragrafo) -> bool:
    """Remove o número digitado, pois a lista multinível passa a gerá-lo."""
    if paragrafo.style.name not in ESTILOS_TITULO:
        return False
    encontrado = RE_TITULO_NUMERADO_LITERAL.match(paragrafo.text.strip())
    if not encontrado:
        return False
    _substituir_texto(paragrafo, encontrado.group(1))
    return True


def _forcar_negrito_titulo(paragrafo) -> None:
    """Evita que um `bold=false` direto da origem anule o estilo em negrito."""
    if paragrafo.style.name not in ESTILOS_TITULO:
        return
    for run in paragrafo.runs:
        _definir_negrito(run._r, True)


def _romano(numero: int) -> str:
    valores = ((1000, "M"), (900, "CM"), (500, "D"), (400, "CD"), (100, "C"), (90, "XC"),
               (50, "L"), (40, "XL"), (10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I"))
    partes = []
    for valor, simbolo in valores:
        while numero >= valor:
            partes.append(simbolo)
            numero -= valor
    return "".join(partes)


def _normalizar_pedido_literal(paragrafo) -> bool:
    if paragrafo.style.name != "FB - Pedido":
        return False
    encontrado = RE_PEDIDO_NUMERADO_LITERAL.match(paragrafo.text)
    if not encontrado:
        return False
    numero, texto = encontrado.groups()
    _substituir_texto(paragrafo, f"{_romano(int(numero))}. {texto}")
    return True


def _converter_tabelas_markdown(doc) -> int:
    paragrafos = list(doc.paragraphs)
    blocos, atual = [], []
    for p in paragrafos:
        if RE_MARKDOWN_TABELA.match(p.text):
            atual.append(p)
        elif atual:
            blocos.append(atual); atual = []
    if atual:
        blocos.append(atual)
    convertidas = 0
    for bloco in blocos:
        linhas = []
        for p in bloco:
            celulas = [c.strip().replace("**", "") for c in p.text.strip().strip("|").split("|")]
            if celulas and all(RE_SEPARADOR_TABELA.match(c.replace(" ", "")) for c in celulas):
                continue
            linhas.append(celulas)
        if not linhas:
            continue
        colunas = max(len(linha) for linha in linhas)
        tabela = doc.add_table(rows=len(linhas), cols=colunas)
        tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
        tabela.autofit = False
        larguras = [Cm(11.5), Cm(4.5)] if colunas == 2 else [Cm(16 / colunas)] * colunas
        for ri, linha in enumerate(linhas):
            for ci in range(colunas):
                celula = tabela.cell(ri, ci)
                celula.width = larguras[min(ci, len(larguras) - 1)]
                celula.text = linha[ci] if ci < len(linha) else ""
                for p in celula.paragraphs:
                    p.style = doc.styles["FB - Tabela"]
        primeiro = bloco[0]._p
        primeiro.addprevious(tabela._tbl)
        for p in bloco:
            p._p.getparent().remove(p._p)
        convertidas += 1
    return convertidas


def _tabela_de_assinatura(tabela) -> bool:
    return len(tabela.rows) == 1 and len(tabela.columns) == 1 and bool(
        re.search(r"\bOAB(?:/\w+)?\b", tabela.cell(0, 0).text, re.IGNORECASE)
    )


def _remover_vazios_antes_da_tabela(tabela) -> int:
    removidos = 0
    anterior = tabela._tbl.getprevious()
    while anterior is not None and anterior.tag == qn("w:p"):
        texto = "".join(anterior.itertext()).strip()
        if texto or anterior.xpath(".//w:drawing"):
            break
        candidato = anterior.getprevious()
        anterior.getparent().remove(anterior)
        removidos += 1
        anterior = candidato
    return removidos


def _remover_borda_inferior(paragrafo) -> None:
    ppr = paragrafo._p.pPr
    if ppr is None:
        return
    bordas = ppr.find(qn("w:pBdr"))
    if bordas is None:
        return
    inferior = bordas.find(qn("w:bottom"))
    if inferior is not None:
        bordas.remove(inferior)
    if len(bordas) == 0:
        ppr.remove(bordas)


def _adicionar_borda_inferior(paragrafo, cor="CCB595") -> None:
    ppr = paragrafo._p.get_or_add_pPr()
    bordas = ppr.find(qn("w:pBdr"))
    if bordas is None:
        bordas = OxmlElement("w:pBdr")
        ppr.append(bordas)
    inferior = OxmlElement("w:bottom")
    inferior.set(qn("w:val"), "single")
    inferior.set(qn("w:sz"), "8")
    inferior.set(qn("w:space"), "4")
    inferior.set(qn("w:color"), cor)
    bordas.append(inferior)


def _finalizar_blocos_de_resumo(paragrafos) -> None:
    for i, paragrafo in enumerate(paragrafos):
        if paragrafo.style.name != "FB - Resumo":
            continue
        _remover_borda_inferior(paragrafo)
        proximo_e_resumo = i + 1 < len(paragrafos) and paragrafos[i + 1].style.name == "FB - Resumo"
        if not proximo_e_resumo:
            _adicionar_borda_inferior(paragrafo)


def _paragrafos_tabelas(tabelas, vistos=None):
    # Manter as referências vivas evita que ids de wrappers temporários de
    # células mescladas sejam reutilizados durante a iteração.
    vistos = [] if vistos is None else vistos
    for tabela in tabelas:
        for linha in tabela.rows:
            for celula in linha.cells:
                if any(celula._tc is tc for tc in vistos):
                    continue
                vistos.append(celula._tc)
                yield from celula.paragraphs
                yield from _paragrafos_tabelas(celula.tables, vistos)


def _importar_estilos(origem: Path, alvo: Path, temporario: Path) -> None:
    src, dst = Document(origem), Document(alvo)
    permitidos = [s for s in src.styles if s.name.startswith("FB - ")]
    estilos_dst = dst.styles.element
    for style in permitidos:
        sid = style.style_id
        existente = estilos_dst.find(f".//w:style[@w:styleId='{sid}']", namespaces=estilos_dst.nsmap)
        if existente is not None:
            estilos_dst.remove(existente)
        estilos_dst.append(copy.deepcopy(style.element))
    dst.save(temporario)


def _limpar_direta(paragrafo, *, preservar_alinhamento: bool = False) -> int:
    alteracoes = 0
    ppr = paragrafo._p.pPr
    if ppr is not None:
        for el in list(ppr):
            nome = el.tag.rsplit("}", 1)[-1]
            if preservar_alinhamento and nome == "jc":
                continue
            if nome in PPR_LIMPAR:
                ppr.remove(el); alteracoes += 1
        marca = ppr.find(qn("w:rPr"))
        if marca is not None:
            for el in list(marca):
                if el.tag.rsplit("}", 1)[-1] in RPR_LIMPAR:
                    marca.remove(el); alteracoes += 1
    for run in paragrafo.runs:
        rpr = run._r.rPr
        if rpr is not None:
            for el in list(rpr):
                if el.tag.rsplit("}", 1)[-1] in RPR_LIMPAR:
                    rpr.remove(el); alteracoes += 1
    return alteracoes


def _intervalos_negrito_qualificacao(texto: str) -> list[tuple[int, int]]:
    """Localiza apenas nomes, entidades, ações e fases processuais."""
    intervalos = [
        (m.start(), m.end())
        for regex in (RE_NOME_ENTIDADE_CAIXA_ALTA, RE_ACAO_FASE_PROCESSUAL)
        for m in regex.finditer(texto)
    ]
    intervalos.sort()
    unidos: list[list[int]] = []
    for inicio, fim in intervalos:
        if unidos and inicio <= unidos[-1][1]:
            unidos[-1][1] = max(unidos[-1][1], fim)
        else:
            unidos.append([inicio, fim])
    return [(inicio, fim) for inicio, fim in unidos]


def _definir_negrito(run_xml, valor: bool) -> None:
    rpr = run_xml.find(qn("w:rPr"))
    if rpr is None and not valor:
        return
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        run_xml.insert(0, rpr)
    for negrito in list(rpr.findall(qn("w:b"))) + list(rpr.findall(qn("w:bCs"))):
        rpr.remove(negrito)
    if valor:
        rpr.append(OxmlElement("w:b"))
    if len(rpr) == 0:
        run_xml.remove(rpr)


def _normalizar_negrito_qualificacao(paragrafo) -> int:
    """Retira negrito global da qualificação e reaplica-o só aos trechos permitidos."""
    intervalos = _intervalos_negrito_qualificacao(paragrafo.text)
    if not paragrafo.runs:
        return 0
    alteracoes = 0
    posicao = 0
    for run in list(paragrafo.runs):
        texto = run.text or ""
        inicio_run, fim_run = posicao, posicao + len(texto)
        posicao = fim_run
        if not texto:
            continue
        cortes = {0, len(texto)}
        for inicio, fim in intervalos:
            if inicio < fim_run and fim > inicio_run:
                cortes.add(max(0, inicio - inicio_run))
                cortes.add(min(len(texto), fim - inicio_run))
        pontos = sorted(cortes)
        partes = [(pontos[i], pontos[i + 1]) for i in range(len(pontos) - 1) if pontos[i] < pontos[i + 1]]
        pai = run._r.getparent()
        indice = pai.index(run._r)
        for parte_inicio, parte_fim in partes:
            trecho = texto[parte_inicio:parte_fim]
            inicio_global = inicio_run + parte_inicio
            fim_global = inicio_run + parte_fim
            deve_negritar = any(inicio_global >= inicio and fim_global <= fim for inicio, fim in intervalos)
            novo_xml = copy.deepcopy(run._r)
            novo_run = Run(novo_xml, paragrafo)
            novo_run.text = trecho
            _definir_negrito(novo_xml, deve_negritar)
            pai.insert(indice, novo_xml)
            indice += 1
        pai.remove(run._r)
        alteracoes += 1
    return alteracoes


def _remover_negrito_direto(paragrafo) -> int:
    """Remove negrito aplicado diretamente em estilos regulares de texto."""
    alteracoes = 0
    for run in paragrafo.runs:
        rpr = run._r.rPr
        if rpr is None:
            continue
        havia_negrito = bool(rpr.find(qn("w:b")) is not None or rpr.find(qn("w:bCs")) is not None)
        if havia_negrito:
            _definir_negrito(run._r, False)
            alteracoes += 1
    return alteracoes


def _imagem_de_assinatura(paragrafo, indice: int, paragrafos, mapa: dict) -> bool:
    """Identifica imagem isolada imediatamente antes do bloco de assinatura."""
    if paragrafo.text.strip() or not paragrafo._p.xpath(".//w:drawing"):
        return False
    textos_vistos = 0
    for proximo_indice in range(indice + 1, len(paragrafos)):
        proximo = paragrafos[proximo_indice]
        if not proximo.text.strip():
            continue
        textos_vistos += 1
        item = mapa.get(proximo_indice + 1, {})
        estilo = item.get("estilo_destino") or (proximo.style.name if proximo.style else "")
        if estilo == "FB - Assinatura" or re.search(r"\b(?:OAB|advogad[oa]|procurador[a]?)\b", proximo.text, re.IGNORECASE):
            return True
        if textos_vistos >= 3:
            return False
    return False


def _aplicar_numpr(paragrafo, numpr) -> None:
    if numpr is None:
        return
    ppr = paragrafo._p.get_or_add_pPr()
    existente = ppr.find(qn("w:numPr"))
    if existente is not None:
        ppr.remove(existente)
    ppr.insert(0, numpr)


def _normalizar_numeracao_pedidos(doc) -> None:
    """Formata os itens do pedido em I, II, III... com texto a 3 cm."""
    numeracao = doc.part.numbering_part.element
    num_para_abstrato = {}
    for num in numeracao.findall(qn("w:num")):
        abstrato = num.find(qn("w:abstractNumId"))
        if abstrato is not None:
            num_para_abstrato[num.get(qn("w:numId"))] = abstrato.get(qn("w:val"))

    niveis_usados = set()
    for p in doc.paragraphs:
        if p.style.name != "FB - Pedido" or p._p.pPr is None:
            continue
        numpr = p._p.pPr.find(qn("w:numPr"))
        if numpr is None:
            continue
        numid = numpr.find(qn("w:numId"))
        ilvl = numpr.find(qn("w:ilvl"))
        if numid is not None:
            niveis_usados.add((numid.get(qn("w:val")), ilvl.get(qn("w:val")) if ilvl is not None else "0"))

    for numid, nivel in niveis_usados:
        abstrato_id = num_para_abstrato.get(numid)
        abstrato = next((x for x in numeracao.findall(qn("w:abstractNum")) if x.get(qn("w:abstractNumId")) == abstrato_id), None)
        if abstrato is None:
            continue
        lvl = next((x for x in abstrato.findall(qn("w:lvl")) if x.get(qn("w:ilvl")) == nivel), None)
        if lvl is None:
            continue
        formato = lvl.find(qn("w:numFmt"))
        if formato is None:
            formato = OxmlElement("w:numFmt")
            lvl.insert(1, formato)
        formato.set(qn("w:val"), "upperRoman")
        ppr = lvl.find(qn("w:pPr"))
        if ppr is None:
            ppr = OxmlElement("w:pPr")
            lvl.append(ppr)
        ind = ppr.find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            ppr.append(ind)
        ind.set(qn("w:left"), "1701")
        ind.set(qn("w:hanging"), "425")
        rpr = lvl.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            lvl.append(rpr)
        fontes = rpr.find(qn("w:rFonts"))
        if fontes is None:
            fontes = OxmlElement("w:rFonts")
            rpr.insert(0, fontes)
        for atributo in ("ascii", "hAnsi", "eastAsia", "cs"):
            fontes.set(qn(f"w:{atributo}"), "Segoe UI Variable Text")


def _estabilizar_papel_timbrado(caminho: Path) -> None:
    """Evita o desaparecimento intermitente de fundos timbrados opacos."""
    ns = {
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
        "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
    }
    with zipfile.ZipFile(caminho) as origem:
        itens = [(info, origem.read(info.filename)) for info in origem.infolist()]
    dados = {info.filename: conteudo for info, conteudo in itens}
    imagens_para_converter = set()
    ancoras_por_imagem = {}
    for nome in [n for n in dados if n.startswith("word/header") and n.endswith(".xml")]:
        rels_nome = f"word/_rels/{Path(nome).name}.rels"
        if rels_nome not in dados:
            continue
        raiz = etree.fromstring(dados[nome])
        rels = etree.fromstring(dados[rels_nome])
        destinos = {x.get("Id"): x.get("Target") for x in rels.findall("pr:Relationship", namespaces=ns)}
        for ancora in raiz.xpath(".//wp:anchor[@behindDoc='1']", namespaces=ns):
            extensao = ancora.find("wp:extent", namespaces=ns)
            blip = ancora.find(".//a:blip", namespaces=ns)
            if extensao is None or blip is None or int(extensao.get("cy", "0")) < 8_000_000:
                continue
            alvo = destinos.get(blip.get(f"{{{ns['r']}}}embed"))
            if not alvo:
                continue
            midia = posixpath.normpath(posixpath.join(posixpath.dirname(nome), alvo))
            if midia in dados:
                imagens_para_converter.add(midia)
                ancoras_por_imagem.setdefault(midia, []).append((nome, ancora, raiz))

    alterados = set()
    for midia in imagens_para_converter:
        try:
            imagem = Image.open(io.BytesIO(dados[midia])).convert("RGBA")
            cinza_invertido = ImageOps.invert(ImageOps.grayscale(imagem))
            alfa = cinza_invertido.point(lambda valor: min(255, valor * 8))
            imagem.putalpha(alfa)
            buffer = io.BytesIO()
            imagem.save(buffer, format="PNG")
            dados[midia] = buffer.getvalue()
        except Exception:
            continue
        for nome, ancora, raiz in ancoras_por_imagem[midia]:
            ancora.set("behindDoc", "0")
            dados[nome] = etree.tostring(raiz, xml_declaration=True, encoding="UTF-8", standalone=True)
            alterados.add(nome)

    if not alterados:
        return
    with tempfile.NamedTemporaryFile(suffix=".docx", dir=caminho.parent, delete=False) as arquivo_tmp:
        temporario = Path(arquivo_tmp.name)
    try:
        with zipfile.ZipFile(temporario, "w") as destino:
            for info, conteudo in itens:
                destino.writestr(info, dados[info.filename])
        temporario.replace(caminho)
    finally:
        if temporario.exists():
            temporario.unlink()


def aplicar(entrada: Path, modelo: Path, classificacao: dict, configuracao: dict, saida: Path) -> dict:
    permitidos = set(configuracao["estilos_permitidos"])
    padrao = configuracao["estilo_padrao"]
    mapa = {x["id"]: x for x in classificacao["classificacoes"]}
    saida.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        estagiado = Path(tmp) / "com_estilos.docx"
        _importar_estilos(modelo, entrada, estagiado)
        doc = Document(estagiado)
        num_id_titulos = criar_lista_multinivel_titulos(doc, vincular_estilos=True)
        contagem, alertas, limpezas, imagens_assinatura_centralizadas, negritos_corpo_removidos = {}, [], 0, 0, 0
        imagens_gerais_centralizadas = 0
        normalizacoes_markdown = titulos_literais_ajustados = pedidos_literais_ajustados = 0
        paragrafos = list(doc.paragraphs)
        for i, p in enumerate(doc.paragraphs, 1):
            item = mapa.get(i, {})
            estilo = item.get("estilo_destino", padrao)
            if estilo not in permitidos:
                alertas.append({"paragrafo": i, "alerta": f"Estilo não autorizado: {estilo}; aplicado {padrao}"})
                estilo = padrao
            if item.get("confianca", 0) < 0.80 and p.text.strip():
                alertas.append({"paragrafo": i, "alerta": "Baixa confiança; aplicado estilo padrão provisório"})
                estilo = padrao
            imagem_isolada = not p.text.strip() and bool(p._p.xpath(".//w:drawing"))
            imagem_assinatura = (
                configuracao.get("regras", {}).get("FB - Assinatura", {}).get("centralizar_imagem", False)
                and _imagem_de_assinatura(p, i - 1, paragrafos, mapa)
            )
            if imagem_assinatura:
                estilo = "FB - Assinatura"
                imagens_assinatura_centralizadas += 1
            numpr = criar_numpr(num_id_titulos, ESTILOS_TITULO[estilo]) if estilo in ESTILOS_TITULO else None
            normalizacoes_markdown += int(_normalizar_markdown_paragrafo(p, estilo))
            limpezas += _limpar_direta(p, preservar_alinhamento=imagem_isolada)
            p.style = doc.styles[estilo]
            if imagem_isolada:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if not imagem_assinatura:
                    imagens_gerais_centralizadas += 1
            if estilo in configuracao.get("regras", {}).get("negrito_direto_remover_estilos", []):
                negritos_corpo_removidos += _remover_negrito_direto(p)
            _aplicar_numpr(p, numpr)
            titulos_literais_ajustados += int(_normalizar_titulo_literal(p))
            _forcar_negrito_titulo(p)
            contagem[estilo] = contagem.get(estilo, 0) + 1
        tabelas_markdown_convertidas = _converter_tabelas_markdown(doc)
        assinaturas_em_tabela = vazios_antes_assinatura_removidos = 0
        for tabela in doc.tables:
            assinatura = _tabela_de_assinatura(tabela)
            if assinatura:
                vazios_antes_assinatura_removidos += _remover_vazios_antes_da_tabela(tabela)
            for p in _paragrafos_tabelas([tabela]):
                limpezas += _limpar_direta(p)
                estilo_tabela = "FB - Assinatura" if assinatura else "FB - Tabela"
                p.style = doc.styles[estilo_tabela]
                if assinatura:
                    assinaturas_em_tabela += 1
                contagem[estilo_tabela] = contagem.get(estilo_tabela, 0) + 1
        # Cabeçalhos e rodapés não são sobrepostos: podem conter papel
        # timbrado, campos de paginação e objetos flutuantes posicionados.
        negritos_qualificacao = 0
        for p in doc.paragraphs:
            if p.style.name == "FB - Qualificação das Partes":
                negritos_qualificacao += _normalizar_negrito_qualificacao(p)
        _normalizar_numeracao_pedidos(doc)
        _finalizar_blocos_de_resumo(doc.paragraphs)
        doc.save(saida)
        _estabilizar_papel_timbrado(saida)
    return {
        "arquivo": str(saida),
        "contagem_estilos": contagem,
        "formatacoes_diretas_removidas": limpezas,
        "negritos_qualificacao_ajustados": negritos_qualificacao,
        "imagens_assinatura_centralizadas": imagens_assinatura_centralizadas,
        "imagens_gerais_centralizadas": imagens_gerais_centralizadas,
        "negritos_corpo_removidos": negritos_corpo_removidos,
        "normalizacoes_markdown": normalizacoes_markdown,
        "titulos_literais_ajustados": titulos_literais_ajustados,
        "pedidos_literais_ajustados": pedidos_literais_ajustados,
        "tabelas_markdown_convertidas": tabelas_markdown_convertidas,
        "assinaturas_em_tabela": assinaturas_em_tabela,
        "vazios_antes_assinatura_removidos": vazios_antes_assinatura_removidos,
        "alertas": alertas,
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("entrada", type=Path); ap.add_argument("classificacao", type=Path)
    ap.add_argument("--modelo", type=Path, default=ROOT / "modelos" / "modelo_oficial.docx")
    ap.add_argument("--config", type=Path, default=ROOT / "modelos" / "estilos.json")
    ap.add_argument("--saida", type=Path, default=ROOT / "saida" / "peticao_formatada.docx")
    args = ap.parse_args()
    print(aplicar(args.entrada, args.modelo, carregar_json(args.classificacao), carregar_json(args.config), args.saida))


if __name__ == "__main__": main()
