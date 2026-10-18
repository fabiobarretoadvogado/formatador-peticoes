from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

try:
    from .comum import ROOT
    from .configurar_visibilidade_estilos import ajustar_arquivo
    from .numeracao_word import criar_lista_multinivel_titulos
except ImportError:
    from comum import ROOT
    from configurar_visibilidade_estilos import ajustar_arquivo
    from numeracao_word import criar_lista_multinivel_titulos


def _fonte(style, nome="Segoe UI Variable Text", tamanho=12, negrito=False, italico=False):
    style.font.name = nome
    style.font.size = Pt(tamanho)
    style.font.bold = negrito
    style.font.italic = italico
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for atributo in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{atributo}"), nome)


def _borda_inferior(style, cor="CCB595"):
    """Insere uma linha inferior discreta no parágrafo do estilo."""
    ppr = style.element.get_or_add_pPr()
    bordas = ppr.find(qn("w:pBdr"))
    if bordas is None:
        bordas = OxmlElement("w:pBdr")
        ppr.append(bordas)
    inferior = bordas.find(qn("w:bottom"))
    if inferior is None:
        inferior = OxmlElement("w:bottom")
        bordas.append(inferior)
    inferior.set(qn("w:val"), "single")
    inferior.set(qn("w:sz"), "8")
    inferior.set(qn("w:space"), "4")
    inferior.set(qn("w:color"), cor)


def criar_modelo(destino: Path | None = None) -> Path:
    destino = destino or ROOT / "modelos" / "modelo_oficial.docx"
    destino.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(3.0)
    sec.right_margin = Cm(2.0)

    specs = {
        "FB - Endereçamento": (WD_ALIGN_PARAGRAPH.LEFT, 0, 0, 12, 12, 12, True, False),
        "FB - Identificação do Processo": (WD_ALIGN_PARAGRAPH.LEFT, 0, 0, 12, 12, 12, True, False),
        "FB - Resumo": (WD_ALIGN_PARAGRAPH.JUSTIFY, 9.0, 0, 12, 12, 10, False, False),
        "FB - Qualificação das Partes": (WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 1.0, 12, 12, 12, False, False),
        "FB - Identificação da Ação": (WD_ALIGN_PARAGRAPH.CENTER, 0, 0, 12, 12, 12, True, False),
        "FB - Título Principal": (WD_ALIGN_PARAGRAPH.LEFT, 0.5, -0.5, 36, 12, 12, True, False),
        "FB - Título Secundário": (WD_ALIGN_PARAGRAPH.LEFT, 1.0, 1.0, 36, 12, 12, True, False),
        "FB - Corpo": (WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 1.0, 12, 12, 12, False, False),
        "FB - Corpo com Recuo": (WD_ALIGN_PARAGRAPH.JUSTIFY, 1.25, 1.0, 12, 12, 12, False, False),
        "FB - Tabela": (WD_ALIGN_PARAGRAPH.LEFT, 0, 0, 0, 0, 12, False, False),
        "FB - Descrição": (WD_ALIGN_PARAGRAPH.LEFT, 2.0, 0, 6, 6, 8, False, False),
        "FB - Citação": (WD_ALIGN_PARAGRAPH.JUSTIFY, 3.0, 0, 12, 12, 12, False, False),
        "FB - Lista": (WD_ALIGN_PARAGRAPH.JUSTIFY, 3.0, 0, 12, 12, 12, False, False),
        "FB - Pedido": (WD_ALIGN_PARAGRAPH.JUSTIFY, 3.0, 0, 12, 12, 12, False, False),
        "FB - Fechamento": (WD_ALIGN_PARAGRAPH.CENTER, 0, 0, 0, 0, 12, False, False),
        "FB - Assinatura": (WD_ALIGN_PARAGRAPH.CENTER, 0, 0, 0, 0, 12, True, False),
        "FB - Nota de Rodapé": (WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 0, 0, 0, 10, False, False),
        "FB - Rodapé": (WD_ALIGN_PARAGRAPH.CENTER, 0, 0, 0, 0, 10, False, False),
    }
    for nome, (alinhamento, esquerdo, primeira, antes, depois, tamanho, negrito, italico) in specs.items():
        style = doc.styles.add_style(nome, WD_STYLE_TYPE.PARAGRAPH)
        _fonte(style, tamanho=tamanho, negrito=negrito, italico=italico)
        if nome in {"FB - Endereçamento", "FB - Identificação do Processo"}:
            style.font.cs_bold = True
            style.font.all_caps = True
            style.font.small_caps = False
        pf = style.paragraph_format
        pf.alignment = alinhamento
        pf.left_indent = Cm(esquerdo)
        pf.right_indent = Cm(2.0) if nome == "FB - Descrição" else None
        pf.first_line_indent = Cm(primeira)
        pf.space_before = Pt(antes)
        pf.space_after = Pt(depois)
        pf.line_spacing = 1.25 if nome not in {"FB - Tabela", "FB - Descrição", "FB - Nota de Rodapé", "FB - Rodapé", "FB - Assinatura", "FB - Fechamento"} else 1.0
        pf.widow_control = False
        pf.keep_with_next = nome in {"FB - Título Principal", "FB - Título Secundário", "FB - Identificação da Ação", "FB - Endereçamento"}

    criar_lista_multinivel_titulos(doc, vincular_estilos=True)

    p = doc.add_paragraph()
    p.style = doc.styles["FB - Título Principal"]
    p.add_run("MODELO OFICIAL DE ESTILOS")
    p_sub = doc.add_paragraph()
    p_sub.style = doc.styles["FB - Título Secundário"]
    p_sub.add_run("SUBTÍTULO DE EXEMPLO")
    p2 = doc.add_paragraph("Arquivo-base: substitua ou ajuste no Microsoft Word após a validação visual do escritório.")
    p2.style = doc.styles["FB - Corpo"]
    doc.save(destino)

    # O modelo continua contendo os estilos internos necessários ao Word, mas
    # o painel de estilos mostra somente os estilos autorais do escritório.
    temporario = destino.with_name(f".{destino.stem}.estilos_visiveis.docx")
    ajustar_arquivo(destino, temporario, specs)
    temporario.replace(destino)
    return destino


if __name__ == "__main__":
    print(criar_modelo())
