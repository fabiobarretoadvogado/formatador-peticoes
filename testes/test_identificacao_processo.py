import json
import tempfile
import unittest
from pathlib import Path

from scripts.classificar_regras import classificar, classificar_um


NUMERO = "0001234-56.2025.8.26.0001"


class IdentificacaoProcessoTests(unittest.TestCase):
    def paragrafo(self, texto, **extras):
        return {"id": 1, "texto": texto, "posicao_relativa": 0.05, **extras}

    def test_numero_unico_e_identificacao_sem_revisao(self):
        texto = "Número único: " + NUMERO
        item = classificar_um(self.paragrafo(texto), "Processo nº 202500123456", "Autor: PARTE AUTORA")
        self.assertEqual(item["estilo_destino"], "FB - Identificação do Processo")
        self.assertEqual(item["confianca"], 0.99)
        self.assertFalse(item["necessita_analise_semantica"])
        self.assertEqual(item["texto"], texto)

    def test_rotulo_admite_caixa_acentos_e_espacos_variados(self):
        for rotulo in ("NÚMERO ÚNICO:", "Numero unico :", "Número   único:\t"):
            with self.subTest(rotulo=rotulo):
                item = classificar_um(self.paragrafo(rotulo + " " + NUMERO))
                self.assertEqual(item["estilo_destino"], "FB - Identificação do Processo")

    def test_outros_numeros_e_referencias_narrativas_nao_sao_identificacao(self):
        for texto in (
            "Número único: 123456",
            "Número único: 123.456.789-00",
            "Número único: 000123-56.2025.8.26.0001",
            "Número único: 0001234-56.2025.8.26.000",
            "Número único: " + NUMERO + "0",
            "O número único: " + NUMERO + " identifica a demanda.",
            "Número único: " + NUMERO + " foi informado pela parte.",
            "Nota fiscal: " + NUMERO,
        ):
            with self.subTest(texto=texto):
                item = classificar_um(self.paragrafo(texto))
                self.assertNotEqual(item["estilo_destino"], "FB - Identificação do Processo")

    def test_rotulos_anteriores_continuam_reconhecidos(self):
        for texto in ("Processo nº 202500123456", "PROCESSO Nº " + NUMERO, "Autos: " + NUMERO, "Proc. " + NUMERO):
            with self.subTest(texto=texto):
                self.assertEqual(classificar_um(self.paragrafo(texto))["estilo_destino"], "FB - Identificação do Processo")

    def test_docx_aplica_identificacao_e_preserva_numero_e_original(self):
        from docx import Document
        from scripts.aplicar_estilos import aplicar
        from scripts.extrair_documento import extrair

        root = Path(__file__).resolve().parents[1]
        with tempfile.TemporaryDirectory() as temp:
            entrada, saida = Path(temp) / "entrada.docx", Path(temp) / "saida.docx"
            doc = Document()
            doc.add_paragraph("AO JUÍZO DA VARA CÍVEL")
            doc.add_paragraph("Processo nº 202500123456")
            texto = "Número único: " + NUMERO
            doc.add_paragraph(texto)
            doc.add_paragraph("Autor: PARTE AUTORA")
            doc.save(entrada)
            original = entrada.read_bytes()
            config = json.loads((root / "modelos/estilos.json").read_text(encoding="utf-8"))
            regras = classificar(extrair(entrada))
            item = regras["classificacoes"][2]
            self.assertFalse(item["necessita_analise_semantica"])
            aplicar(entrada, root / "modelos/modelo_oficial.docx", regras, config, saida)
            p = Document(saida).paragraphs[2]
            self.assertEqual(p.style.name, "FB - Identificação do Processo")
            self.assertEqual(p.text, texto)
            self.assertEqual(entrada.read_bytes(), original)
