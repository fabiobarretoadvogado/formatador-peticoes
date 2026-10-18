import json
import tempfile
import unittest
from pathlib import Path

from scripts.classificar_regras import classificar, classificar_um


class TitulosEnderecamentoTests(unittest.TestCase):
    def paragrafo(self, texto, **extras):
        return {"id": 1, "texto": texto, "posicao_relativa": 0.5, **extras}

    def test_sintese_numerada_e_titulo_principal_em_qualquer_posicao(self):
        for posicao in (0.05, 0.5, 0.9):
            with self.subTest(posicao=posicao):
                texto = "I. SÍNTESE DA CONTROVÉRSIA"
                item = classificar_um(self.paragrafo(texto, posicao_relativa=posicao))
                self.assertEqual(item["estilo_destino"], "FB - Título Principal")
                self.assertFalse(item["necessita_analise_semantica"])
                self.assertEqual(item["texto"], texto)

    def test_titulos_enumerados_respeitam_hierarquia(self):
        for texto, estilo in (
            ("II. DOS FATOS", "FB - Título Principal"),
            ("IV – DA TUTELA DE URGÊNCIA", "FB - Título Principal"),
            ("1. Dos fundamentos jurídicos", "FB - Título Principal"),
            ("1.1. DA PROVA DOCUMENTAL", "FB - Título Secundário"),
            ("2.3 Da responsabilidade civil", "FB - Título Secundário"),
        ):
            with self.subTest(texto=texto):
                self.assertEqual(classificar_um(self.paragrafo(texto))["estilo_destino"], estilo)
        item = classificar_um(self.paragrafo("I. SÍNTESE DA CONTROVÉRSIA", estilo_original="Heading 2"))
        self.assertEqual(item["estilo_destino"], "FB - Título Secundário")

    def test_listas_e_frases_numeradas_nao_viram_titulos(self):
        for texto in (
            "I. A procedência integral do pedido;",
            "1. Cópia do contrato;",
            "II. PAGAMENTO DO DÉBITO",
            "1. Dos documentos juntados consta a assinatura do autor.",
            "I. DOS FATOS NARRADOS, CONCLUI-SE A NECESSIDADE DA PROVA.",
        ):
            with self.subTest(texto=texto):
                self.assertEqual(classificar_um(self.paragrafo(texto))["estilo_destino"], "FB - Lista")

    def test_vocativo_ao_tribunal_e_enderecamento(self):
        for texto in (
            "EGRÉGIO SUPERIOR TRIBUNAL DE JUSTIÇA",
            "Egrégio Tribunal de Justiça do Estado de Sergipe",
            "AO EGRÉGIO TRIBUNAL REGIONAL FEDERAL DA 5ª REGIÃO",
            "COLENDA TURMA",
            "À EGRÉGIA CÂMARA CÍVEL",
        ):
            with self.subTest(texto=texto):
                item = classificar_um(self.paragrafo(texto, estilo_original="Heading 1"))
                self.assertEqual(item["estilo_destino"], "FB - Endereçamento")
                self.assertFalse(item["necessita_analise_semantica"])

    def test_mencoes_e_citacoes_a_tribunais_preservam_funcao(self):
        for texto in (
            "O Egrégio Superior Tribunal de Justiça decidiu sobre a matéria.",
            "Conforme decidiu o Egrégio Tribunal de Justiça, aplica-se o entendimento.",
            "Egrégio Superior Tribunal de Justiça decidiu pela improcedência do pedido.",
        ):
            with self.subTest(texto=texto):
                self.assertNotEqual(classificar_um(self.paragrafo(texto))["estilo_destino"], "FB - Endereçamento")
        for texto in ("EGRÉGIO SUPERIOR TRIBUNAL DE JUSTIÇA", "I. SÍNTESE DA CONTROVÉRSIA"):
            item = classificar_um(self.paragrafo(texto, estilo_original="Quote"))
            self.assertEqual(item["estilo_destino"], "FB - Citação")

    def test_titulo_de_pedidos_e_itens_numerados_sao_distintos(self):
        estrutura = {"documento": "amostra.docx", "paragrafos": [
            self.paragrafo("I. DOS FATOS"),
            self.paragrafo("II. DOS PEDIDOS", id=2),
            self.paragrafo("I. A procedência do pedido;", id=3),
            self.paragrafo("II. A produção de provas;", id=4),
        ]}
        self.assertEqual([p["estilo_destino"] for p in classificar(estrutura)["classificacoes"]], [
            "FB - Título Principal", "FB - Título Principal", "FB - Pedido", "FB - Pedido",
        ])

    def test_docx_aplica_enderecamento_e_titulo_com_numeracao_nativa(self):
        from docx import Document
        from docx.oxml.ns import qn
        from scripts.aplicar_estilos import aplicar
        from scripts.extrair_documento import extrair

        root = Path(__file__).resolve().parents[1]
        with tempfile.TemporaryDirectory() as temp:
            entrada, saida = Path(temp) / "entrada.docx", Path(temp) / "saida.docx"
            doc = Document()
            doc.add_paragraph("EGRÉGIO SUPERIOR TRIBUNAL DE JUSTIÇA")
            doc.add_paragraph("I. SÍNTESE DA CONTROVÉRSIA")
            doc.add_paragraph("Texto de exemplo para verificar a classificação estrutural.")
            doc.save(entrada)
            original = entrada.read_bytes()
            regras = classificar(extrair(entrada))
            config = json.loads((root / "modelos/estilos.json").read_text(encoding="utf-8"))
            aplicar(entrada, root / "modelos/modelo_oficial.docx", regras, config, saida)
            final = Document(saida)
            self.assertEqual(final.paragraphs[0].style.name, "FB - Endereçamento")
            self.assertEqual(final.paragraphs[0].text, "EGRÉGIO SUPERIOR TRIBUNAL DE JUSTIÇA")
            titulo = final.paragraphs[1]
            self.assertEqual(titulo.style.name, "FB - Título Principal")
            self.assertEqual(titulo.text, "SÍNTESE DA CONTROVÉRSIA")
            self.assertIsNotNone(titulo._p.pPr.find(qn("w:numPr")))
            self.assertEqual(entrada.read_bytes(), original)
