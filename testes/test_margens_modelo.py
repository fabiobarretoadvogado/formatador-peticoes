import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.shared import Cm

from scripts.aplicar_estilos import aplicar
from scripts.classificar_regras import classificar
from scripts.extrair_documento import extrair


ROOT = Path(__file__).resolve().parents[1]
MODELO = ROOT / "modelos/modelo_oficial.docx"
CONFIG = json.loads((ROOT / "modelos/estilos.json").read_text(encoding="utf-8"))


class MargensModeloTests(unittest.TestCase):
    def processar(self, doc, pasta, modelo=MODELO):
        entrada, saida = pasta / "entrada.docx", pasta / "saida.docx"
        doc.save(entrada)
        original = entrada.read_bytes()
        modelo_original = modelo.read_bytes()
        relatorio = aplicar(entrada, modelo, classificar(extrair(entrada)), CONFIG, saida)
        self.assertEqual(entrada.read_bytes(), original)
        self.assertEqual(modelo.read_bytes(), modelo_original)
        return Document(saida), relatorio, saida

    def test_margens_da_imagem_sao_substituidas_pelas_do_modelo(self):
        doc = Document()
        doc.add_paragraph("Conteúdo original da petição.")
        secao = doc.sections[0]
        secao.top_margin, secao.bottom_margin = Cm(3), Cm(2)
        secao.left_margin, secao.right_margin = Cm(3), Cm(2)
        secao.header_distance = secao.footer_distance = Cm(1.7)
        with tempfile.TemporaryDirectory() as temp:
            resultado, relatorio, _ = self.processar(doc, Path(temp))
            referencia = Document(MODELO).sections[0]
            self.assertEqual(resultado.sections[0]._sectPr.pgMar.attrib, referencia._sectPr.pgMar.attrib)
            self.assertEqual(resultado.paragraphs[0].text, doc.paragraphs[0].text)
            self.assertEqual(len(relatorio["margens_modelo"]), 1)
            registro = relatorio["margens_modelo"][0]
            self.assertEqual(registro["antes_cm"]["esquerda"], 3)
            self.assertEqual(registro["depois_cm"]["esquerda"], 2)
            self.assertAlmostEqual(registro["depois_cm"]["inferior"], 2.75, places=2)

    def test_todas_secoes_recebem_margens_sem_mudar_tamanho_orientacao_ou_quebras(self):
        doc = Document()
        doc.add_paragraph("Primeira seção.")
        secao = doc.add_section(WD_SECTION_START.NEW_PAGE)
        secao.orientation = WD_ORIENT.LANDSCAPE
        secao.page_width, secao.page_height = Cm(29.7), Cm(21)
        secao.left_margin = Cm(3.5)
        doc.add_paragraph("Segunda seção em paisagem.")
        secao = doc.add_section(WD_SECTION_START.CONTINUOUS)
        secao.bottom_margin = Cm(2)
        doc.add_paragraph("Terceira seção contínua.")
        estrutura = [(s.page_width, s.page_height, s.orientation, s.start_type) for s in doc.sections]
        with tempfile.TemporaryDirectory() as temp:
            resultado, _, _ = self.processar(doc, Path(temp))
            self.assertEqual(len(resultado.sections), 3)
            referencia = Document(MODELO).sections[0]._sectPr.pgMar.attrib
            for indice, secao in enumerate(resultado.sections):
                self.assertEqual(secao._sectPr.pgMar.attrib, referencia)
                self.assertEqual((secao.page_width, secao.page_height, secao.orientation, secao.start_type), estrutura[indice])

    def test_margens_sao_lidas_do_modelo_recebido_sem_valores_fixos_no_codigo(self):
        doc = Document()
        doc.add_paragraph("Conteúdo original.")
        with tempfile.TemporaryDirectory() as temp:
            modelo = Path(temp) / "modelo_aprovado.docx"
            personalizado = Document(MODELO)
            referencia = personalizado.sections[0]
            referencia.left_margin, referencia.bottom_margin = Cm(2.3), Cm(2.6)
            referencia.header_distance, referencia.footer_distance = Cm(1.3), Cm(1.4)
            personalizado.save(modelo)
            resultado, _, _ = self.processar(doc, Path(temp), modelo)
            self.assertEqual(resultado.sections[0]._sectPr.pgMar.attrib, referencia._sectPr.pgMar.attrib)

    def test_reprocessar_nao_altera_margens_ja_aplicadas(self):
        doc = Document()
        doc.add_paragraph("Conteúdo original.")
        with tempfile.TemporaryDirectory() as temp:
            _, _, primeira_saida = self.processar(doc, Path(temp))
            segunda_saida = Path(temp) / "segunda.docx"
            relatorio = aplicar(primeira_saida, MODELO, classificar(extrair(primeira_saida)), CONFIG, segunda_saida)
            self.assertEqual(relatorio["margens_modelo"], [])
            self.assertEqual(Document(segunda_saida).sections[0]._sectPr.pgMar.attrib, Document(MODELO).sections[0]._sectPr.pgMar.attrib)
