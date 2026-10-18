import io
import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml.ns import qn
from docx.shared import Cm
from PIL import Image

from scripts.aplicar_estilos import aplicar
from scripts.classificar_regras import classificar
from scripts.extrair_documento import extrair


ROOT = Path(__file__).resolve().parents[1]
MODELO = ROOT / "modelos/modelo_oficial.docx"
CONFIG = json.loads((ROOT / "modelos/estilos.json").read_text(encoding="utf-8"))


class PapelTimbradoTests(unittest.TestCase):
    def processar(self, doc, pasta):
        entrada, saida = pasta / "entrada.docx", pasta / "saida.docx"
        doc.save(entrada)
        original = entrada.read_bytes()
        relatorio = aplicar(entrada, MODELO, classificar(extrair(entrada)), CONFIG, saida)
        self.assertEqual(entrada.read_bytes(), original)
        return Document(saida), relatorio, saida

    def assert_timbre(self, cabecalho):
        imagens = cabecalho._element.xpath(".//a:blip")
        self.assertTrue(imagens)
        for imagem in imagens:
            rel = cabecalho.part.rels[imagem.get(qn("r:embed"))]
            self.assertTrue(rel.target_part.blob)
            with Image.open(io.BytesIO(rel.target_part.blob)) as imagem_timbre:
                self.assertGreater(imagem_timbre.width, 1000)

    def test_word_sem_cabecalho_recebe_timbre_e_mantem_rodape_existente(self):
        doc = Document()
        doc.add_paragraph("Texto original da petição.")
        doc.sections[0].footer.paragraphs[0].text = "Rodapé original com paginação própria"
        modelo_original = MODELO.read_bytes()
        with tempfile.TemporaryDirectory() as temp:
            resultado, relatorio, _ = self.processar(doc, Path(temp))
            self.assert_timbre(resultado.sections[0].header)
            self.assertEqual(resultado.sections[0].footer.paragraphs[0].text, "Rodapé original com paginação própria")
            self.assertEqual(resultado.paragraphs[0].text, "Texto original da petição.")
            self.assertEqual(relatorio["papel_timbrado"], {"cabecalhos_adicionados": 1, "rodapes_adicionados": 0})
        self.assertEqual(MODELO.read_bytes(), modelo_original)

    def test_word_sem_rodape_recebe_campos_de_paginacao_do_modelo(self):
        doc = Document()
        doc.add_paragraph("Texto original.")
        with tempfile.TemporaryDirectory() as temp:
            resultado, relatorio, _ = self.processar(doc, Path(temp))
            self.assert_timbre(resultado.sections[0].header)
            campos = " ".join(resultado.sections[0].footer._element.xpath(".//w:instrText/text()"))
            self.assertIn("PAGE", campos)
            self.assertIn("NUMPAGES", campos)
            self.assertEqual(relatorio["papel_timbrado"]["rodapes_adicionados"], 1)

    def test_cabecalho_e_rodape_existentes_nao_sao_substituidos(self):
        doc = Document()
        secao = doc.sections[0]
        secao.top_margin = Cm(2.1)
        secao.header_distance = Cm(0.8)
        secao.header.paragraphs[0].text = "Cabeçalho próprio"
        secao.footer.paragraphs[0].text = "Rodapé próprio"
        doc.add_paragraph("Texto original.")
        with tempfile.TemporaryDirectory() as temp:
            resultado, relatorio, _ = self.processar(doc, Path(temp))
            self.assertEqual(resultado.sections[0].header.paragraphs[0].text, "Cabeçalho próprio")
            self.assertEqual(resultado.sections[0].footer.paragraphs[0].text, "Rodapé próprio")
            self.assertEqual(resultado.sections[0].top_margin, secao.top_margin)
            self.assertEqual(resultado.sections[0].header_distance, secao.header_distance)
            self.assertEqual(relatorio["papel_timbrado"], {"cabecalhos_adicionados": 0, "rodapes_adicionados": 0})

    def test_timbre_nao_colide_com_imagem_do_corpo(self):
        doc = Document()
        doc.add_paragraph("Texto original.")
        imagem = io.BytesIO()
        Image.new("RGB", (31, 29), "red").save(imagem, format="PNG")
        original = imagem.getvalue()
        doc.add_picture(io.BytesIO(original), width=Cm(2))
        with tempfile.TemporaryDirectory() as temp:
            resultado, _, saida = self.processar(doc, Path(temp))
            self.assert_timbre(resultado.sections[0].header)
            self.assertEqual(len(resultado.inline_shapes), 1)
            blip = resultado.inline_shapes[0]._inline.xpath(".//a:blip")[0]
            self.assertEqual(resultado.part.rels[blip.get(qn("r:embed"))].target_part.blob, original)
            with zipfile.ZipFile(saida) as pacote:
                self.assertEqual(len([n for n in pacote.namelist() if n.startswith("word/media/")]), 2)

    def test_primeira_pagina_e_paginas_pares_recebem_timbre(self):
        doc = Document()
        doc.sections[0].different_first_page_header_footer = True
        doc.settings.odd_and_even_pages_header_footer = True
        doc.add_paragraph("Texto original.")
        with tempfile.TemporaryDirectory() as temp:
            resultado, relatorio, _ = self.processar(doc, Path(temp))
            secao = resultado.sections[0]
            for cabecalho in (secao.header, secao.first_page_header, secao.even_page_header):
                self.assert_timbre(cabecalho)
            self.assertEqual(relatorio["papel_timbrado"]["cabecalhos_adicionados"], 3)
            self.assertTrue(secao.different_first_page_header_footer)
            self.assertTrue(resultado.settings.odd_and_even_pages_header_footer)

    def test_secoes_vinculadas_e_independentes_mantem_timbre(self):
        doc = Document()
        doc.add_paragraph("Primeira seção.")
        doc.add_section(WD_SECTION_START.NEW_PAGE)
        doc.add_paragraph("Segunda seção vinculada.")
        secao = doc.add_section(WD_SECTION_START.NEW_PAGE)
        secao.header.is_linked_to_previous = False
        doc.add_paragraph("Terceira seção independente.")
        with tempfile.TemporaryDirectory() as temp:
            resultado, relatorio, _ = self.processar(doc, Path(temp))
            self.assertEqual(len(resultado.sections), 3)
            for secao in resultado.sections:
                self.assert_timbre(secao.header)
            self.assertTrue(resultado.sections[1].header.is_linked_to_previous)
            self.assertFalse(resultado.sections[2].header.is_linked_to_previous)
            self.assertEqual(relatorio["papel_timbrado"]["cabecalhos_adicionados"], 2)

    def test_reprocessamento_nao_duplica_timbre(self):
        doc = Document()
        doc.add_paragraph("Texto original.")
        with tempfile.TemporaryDirectory() as temp:
            resultado, _, primeira_saida = self.processar(doc, Path(temp))
            segunda_saida = Path(temp) / "reprocessado.docx"
            relatorio = aplicar(primeira_saida, MODELO, classificar(extrair(primeira_saida)), CONFIG, segunda_saida)
            final = Document(segunda_saida)
            self.assertEqual(len(final.sections[0].header._element.xpath(".//a:blip")), 1)
            self.assertEqual(relatorio["papel_timbrado"], {"cabecalhos_adicionados": 0, "rodapes_adicionados": 0})
