import json
import tempfile
import unittest
from pathlib import Path

from scripts.classificar_regras import classificar, classificar_um


TRECHO = (
    "O recurso, contudo, não reúne condições de admissibilidade e, de todo modo, "
    "não demonstra qualquer erro no acórdão recorrido."
)


class CorpoDiscursivoTests(unittest.TestCase):
    def paragrafo(self, texto, **extras):
        return {"id": 13, "texto": texto, "posicao_relativa": 0.4, **extras}

    def test_trecho_da_imagem_e_corpo_sem_revisao_desnecessaria(self):
        item = classificar_um(self.paragrafo(TRECHO))
        self.assertEqual(item["estilo_destino"], "FB - Corpo")
        self.assertGreaterEqual(item["confianca"], 0.95)
        self.assertFalse(item["necessita_analise_semantica"])
        self.assertEqual(item["texto"], TRECHO)

    def test_frases_argumentativas_medias_sao_reconhecidas(self):
        for texto in (
            "A parte autora não comprova o requisito necessário ao acolhimento do pedido.",
            "Por isso, a conclusão decorre dos documentos apresentados nos autos.",
            "Este parágrafo apresenta os fundamentos da pretensão deduzida em juízo.",
        ):
            with self.subTest(texto=texto):
                item = classificar_um(self.paragrafo(texto))
                self.assertEqual(item["estilo_destino"], "FB - Corpo")
                self.assertFalse(item["necessita_analise_semantica"])

    def test_quantidade_de_palavras_sozinha_nao_eleva_confianca(self):
        for texto in (
            "Dados relativos à documentação apresentada pela parte autora para análise da situação processual.",
            "Documentos e elementos relativos à situação da parte autora para apreciação do pedido.",
        ):
            with self.subTest(texto=texto):
                item = classificar_um(self.paragrafo(texto))
                self.assertLess(item["confianca"], 0.95)
                self.assertTrue(item["necessita_analise_semantica"])
        self.assertTrue(classificar_um(self.paragrafo("Texto curto."))["necessita_analise_semantica"])

    def test_citacoes_listas_e_titulos_preservam_prioridade(self):
        for p, estilo in (
            (self.paragrafo(TRECHO, estilo_original="Quote"), "FB - Citação"),
            (self.paragrafo(TRECHO, recuo_esquerdo_cm=3), "FB - Citação"),
            (self.paragrafo("1. " + TRECHO), "FB - Lista"),
            (self.paragrafo(TRECHO, possui_numeracao=True), "FB - Lista"),
            (self.paragrafo(TRECHO, estilo_original="Heading 1"), "FB - Título Principal"),
        ):
            with self.subTest(p=p):
                self.assertEqual(classificar_um(p)["estilo_destino"], estilo)

    def test_rotulo_de_descricao_continua_pendente(self):
        texto = "Descrição: a imagem demonstra os elementos apresentados pela parte autora nos autos."
        self.assertTrue(classificar_um(self.paragrafo(texto))["necessita_analise_semantica"])

    def test_docx_preserva_texto_original_e_aplica_corpo(self):
        from docx import Document
        from scripts.aplicar_estilos import aplicar
        from scripts.extrair_documento import extrair

        root = Path(__file__).resolve().parents[1]
        with tempfile.TemporaryDirectory() as temp:
            entrada, saida = Path(temp) / "entrada.docx", Path(temp) / "saida.docx"
            doc = Document()
            doc.add_paragraph(TRECHO)
            doc.save(entrada)
            original = entrada.read_bytes()
            regras = classificar(extrair(entrada))
            self.assertFalse(regras["ambiguos"])
            config = json.loads((root / "modelos/estilos.json").read_text(encoding="utf-8"))
            aplicar(entrada, root / "modelos/modelo_oficial.docx", regras, config, saida)
            paragrafo = Document(saida).paragraphs[0]
            self.assertEqual(paragrafo.text, TRECHO)
            self.assertEqual(paragrafo.style.name, "FB - Corpo")
            self.assertEqual(entrada.read_bytes(), original)
