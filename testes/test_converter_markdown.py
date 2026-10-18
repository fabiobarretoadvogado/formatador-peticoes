from __future__ import annotations

import json
import shutil
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from PIL import Image

from scripts.comum import carregar_json, salvar_json
from scripts.converter_markdown import (
    PERFIL_MARKDOWN_PADRAO,
    RecursoMarkdownNaoSuportado,
    converter_markdown,
)
from scripts.pipeline import executar


ROOT = Path(__file__).parents[1]


def _workspace(base: Path) -> Path:
    root = base / "execucao"
    modelos = root / "modelos"
    modelos.mkdir(parents=True)
    shutil.copy2(ROOT / "modelos" / "modelo_oficial.docx", modelos / "modelo_oficial.docx")
    shutil.copy2(ROOT / "modelos" / "estilos.json", modelos / "estilos.json")
    return root


def _markdown_completo(pasta: Path) -> Path:
    Image.new("RGB", (120, 60), "white").save(pasta / "figura.png")
    origem = pasta / "peticao.md"
    origem.write_text(
        "## DOS FATOS\n\n"
        "Texto com **negrito** e *itálico*.\n\n"
        "### DA PROVA\n\n"
        "> Trecho citado em bloco.\n\n"
        "1. primeiro item\n2. segundo item\n\n"
        "a) item alfa\nb) item beta\n\n"
        "- item com marcador\n\n"
        "| Campo | Valor |\n| --- | --- |\n| Origem | Markdown |\n\n"
        "![Figura de teste](figura.png)\n",
        encoding="utf-8",
    )
    return origem


class ConverterMarkdownTests(unittest.TestCase):
    def test_perfil_markdown_do_codigo_reflete_as_regras_do_projeto(self):
        config = carregar_json(ROOT / "modelos" / "estilos.json")
        self.assertEqual(config["formatos_entrada_suportados"], ["docx", "markdown"])
        self.assertEqual(config["perfil_markdown"], PERFIL_MARKDOWN_PADRAO)

    def test_conversor_preserva_estrutura_runs_tabela_lista_e_imagem(self):
        with tempfile.TemporaryDirectory() as d:
            base = Path(d)
            origem = _markdown_completo(base)
            saida = base / "intermediario.docx"
            manifesto = base / "manifesto.json"

            dados = converter_markdown(
                origem,
                ROOT / "modelos" / "modelo_oficial.docx",
                saida,
                manifesto,
            )
            doc = Document(saida)
            titulo = next(p for p in doc.paragraphs if p.text == "DOS FATOS")
            subtitulo = next(p for p in doc.paragraphs if p.text == "DA PROVA")
            corpo = next(p for p in doc.paragraphs if p.text.startswith("Texto com"))
            citacao = next(p for p in doc.paragraphs if p.text.startswith("Trecho citado"))

            self.assertEqual(titulo.style.name, "Heading 1")
            self.assertEqual(subtitulo.style.name, "Heading 2")
            self.assertEqual(citacao.style.name, "Quote")
            self.assertIn("negrito", "".join(run.text for run in corpo.runs if run.bold))
            self.assertIn("itálico", "".join(run.text for run in corpo.runs if run.italic))
            itens = [p for p in doc.paragraphs if p.text.startswith(("primeiro", "segundo", "item alfa", "item beta", "item com"))]
            self.assertEqual(len(itens), 5)
            self.assertTrue(all(p._p.pPr.find(qn("w:numPr")) is not None for p in itens))
            self.assertEqual(len(doc.tables), 1)
            self.assertEqual(doc.tables[0].cell(1, 1).text, "Markdown")
            imagens = [p for p in doc.paragraphs if p._p.xpath(".//w:drawing")]
            self.assertEqual(len(imagens), 1)
            self.assertEqual(imagens[0].alignment, WD_ALIGN_PARAGRAPH.CENTER)
            self.assertEqual(dados["contagem"]["listas_alfabeticas"], 1)
            self.assertEqual(dados["listas_alfabeticas_preprocessadas"], 2)
            self.assertEqual(dados["contagem"]["imagens"], 1)
            self.assertTrue(dados["auditoria_conteudo"]["aprovada"])
            self.assertEqual(json.loads(manifesto.read_text(encoding="utf-8"))["sha256_origem"], dados["sha256_origem"])

    def test_pipeline_aceita_markdown_e_registra_intermediario(self):
        with tempfile.TemporaryDirectory() as d:
            base = Path(d)
            origem = _markdown_completo(base)
            original = origem.read_bytes()
            resultado = executar(origem, skip_pdf=True, work_root=_workspace(base))

            self.assertEqual(origem.read_bytes(), original)
            self.assertEqual(resultado["formato_original"], "markdown")
            self.assertTrue(Path(resultado["docx_intermediario"]).is_file())
            self.assertTrue(Path(resultado["manifesto_ingestao"]).is_file())
            self.assertEqual(Path(resultado["backup"]).suffix, ".md")
            relatorio = Path(resultado["relatorio"]).read_text(encoding="utf-8")
            self.assertIn("Formato original: **Markdown (.md)**", relatorio)
            self.assertIn("DOCX intermediário utilizado", relatorio)

            final = Document(resultado["docx"])
            titulo = next(p for p in final.paragraphs if p.text == "DOS FATOS")
            corpo = next(p for p in final.paragraphs if p.text.startswith("Texto com"))
            self.assertEqual(titulo.style.name, "FB - Título Principal")
            self.assertTrue(all(not bool(run.bold) for run in corpo.runs))
            self.assertIn("itálico", "".join(run.text for run in corpo.runs if run.italic))
            self.assertEqual(len(final.tables), 1)
            self.assertTrue(all(p.style.name == "FB - Tabela" for p in final.tables[0].cell(0, 0).paragraphs))
            imagem = next(p for p in final.paragraphs if p._p.xpath(".//w:drawing"))
            self.assertEqual(imagem.alignment, WD_ALIGN_PARAGRAPH.CENTER)

    def test_pipeline_respeita_formatos_autorizados_na_configuracao(self):
        with tempfile.TemporaryDirectory() as d:
            base = Path(d)
            origem = _markdown_completo(base)
            workspace = _workspace(base)
            config_path = workspace / "modelos" / "estilos.json"
            config = carregar_json(config_path)
            config["formatos_entrada_suportados"] = ["docx"]
            salvar_json(config_path, config)
            with self.assertRaisesRegex(ValueError, "não está autorizado"):
                executar(origem, skip_pdf=True, work_root=workspace)

    def test_rejeita_recurso_fora_do_perfil_sem_descartar_silenciosamente(self):
        with tempfile.TemporaryDirectory() as d:
            base = Path(d)
            origem = base / "invalido.md"
            origem.write_text("# Título não suportado\n", encoding="utf-8")
            with self.assertRaisesRegex(RecursoMarkdownNaoSuportado, "use apenas ## e ###"):
                converter_markdown(
                    origem,
                    ROOT / "modelos" / "modelo_oficial.docx",
                    base / "saida.docx",
                    base / "manifesto.json",
                )

    def test_rejeita_imagem_fora_da_pasta_da_origem(self):
        with tempfile.TemporaryDirectory() as d:
            base = Path(d)
            origem = base / "invalido.md"
            origem.write_text("![fora](../fora.png)\n", encoding="utf-8")
            with self.assertRaisesRegex(RecursoMarkdownNaoSuportado, "dentro da pasta"):
                converter_markdown(
                    origem,
                    ROOT / "modelos" / "modelo_oficial.docx",
                    base / "saida.docx",
                    base / "manifesto.json",
                )


if __name__ == "__main__":
    unittest.main()
