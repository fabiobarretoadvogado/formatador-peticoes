from __future__ import annotations

import argparse
import hashlib
import json
import os
import subprocess
import tempfile
import zipfile
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.shared import Cm, Inches
from lxml import etree
from PIL import Image, ImageChops, ImageDraw


NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}
W = f"{{{NS['w']}}}"
R = f"{{{NS['r']}}}"


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    digest.update(path.read_bytes())
    return digest.hexdigest()


def _xml(data: bytes) -> etree._Element:
    return etree.fromstring(data)


def _create_word_control(path: Path, image_path: Path) -> dict[str, object]:
    image = Image.new("RGB", (240, 80), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((4, 4, 235, 75), outline="#0057a8", width=5)
    draw.text((18, 27), "IMAGEM DE CONTROLE WORD", fill="#0057a8")
    image.save(image_path)

    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Cm(1.7)
    section.right_margin = Cm(2.1)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.3)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(1.1)
    section.header.paragraphs[0].text = "CABEÇALHO DE CONTROLE — ENTRADA WORD"
    section.footer.paragraphs[0].text = "RODAPÉ DE CONTROLE — ENTRADA WORD"
    document.add_heading("DOS FATOS", level=1)
    document.add_paragraph(
        "Texto jurídico de controle que deve permanecer integralmente no documento processado."
    )
    document.add_picture(str(image_path), width=Cm(6))
    document.add_paragraph("Fecho de controle da petição em Word.")
    document.save(path)
    return _word_snapshot(path)


def _word_snapshot(path: Path) -> dict[str, object]:
    document = Document(path)
    section = document.sections[0]
    return {
        "page_width": section.page_width,
        "page_height": section.page_height,
        "top_margin": section.top_margin,
        "right_margin": section.right_margin,
        "bottom_margin": section.bottom_margin,
        "left_margin": section.left_margin,
        "header_distance": section.header_distance,
        "footer_distance": section.footer_distance,
        "header_text": "\n".join(p.text for p in section.header.paragraphs),
        "footer_text": "\n".join(p.text for p in section.footer.paragraphs),
        "body_text": "\n".join(p.text for p in document.paragraphs),
        "inline_images": len(document.inline_shapes),
    }


def _run(exe: Path, input_path: Path, output_root: Path, local_app_data: Path) -> Path:
    output_root.mkdir(parents=True, exist_ok=True)
    environment = os.environ.copy()
    environment["LOCALAPPDATA"] = str(local_app_data)
    result = subprocess.run(
        [
            str(exe),
            str(input_path),
            "--saida",
            str(output_root),
            "--sem-pdf",
            "--silencioso",
        ],
        cwd=exe.parent,
        env=environment,
        capture_output=True,
        text=True,
        timeout=180,
        check=False,
    )
    if result.returncode != 0:
        raise RuntimeError(
            f"O executável falhou para {input_path.name} (código {result.returncode}).\n"
            f"STDOUT:\n{result.stdout}\nSTDERR:\n{result.stderr}"
        )
    outputs = list(output_root.rglob("*_formatada.docx"))
    if len(outputs) != 1:
        raise RuntimeError(f"Esperava-se uma saída DOCX em {output_root}; encontradas: {outputs}")
    return outputs[0]


def _assert_word_preserved(before: dict[str, object], output: Path, model: Path) -> dict[str, object]:
    after = _word_snapshot(output)
    expected = dict(before)
    reference = Document(model).sections[0]
    for key in ("top_margin", "right_margin", "bottom_margin", "left_margin", "header_distance", "footer_distance"):
        expected[key] = getattr(reference, key)
    differing = {key: (expected[key], after[key]) for key in expected if expected[key] != after[key]}
    if differing:
        raise RuntimeError(f"A entrada Word teve elementos fora dos estilos alterados: {differing}")

    with zipfile.ZipFile(output) as package:
        document_xml = _xml(package.read("word/document.xml"))
        rels = _xml(package.read("word/_rels/document.xml.rels"))
        section = document_xml.find("w:body/w:sectPr", namespaces=NS)
        if section is None:
            raise RuntimeError("A saída Word perdeu as propriedades de seção.")
        footer_ref = section.find("w:footerReference", namespaces=NS)
        if footer_ref is None:
            raise RuntimeError("A saída Word perdeu o rodapé de controle.")
        footer_id = footer_ref.get(f"{R}id")
        footer_rel = next(
            rel for rel in rels.findall(f"{{{NS['pr']}}}Relationship") if rel.get("Id") == footer_id
        )
        footer_name = "word/" + footer_rel.get("Target", "").lstrip("/")
        footer = _xml(package.read(footer_name))
        fields = " ".join(footer.xpath(".//w:instrText/text()", namespaces=NS))
        if "PAGE" in fields or "NUMPAGES" in fields:
            raise RuntimeError("A paginação do modelo foi inserida indevidamente na entrada Word.")

    return after


def _relationship_target(rels: etree._Element, rel_id: str) -> str:
    relationship = next(
        rel for rel in rels.findall(f"{{{NS['pr']}}}Relationship") if rel.get("Id") == rel_id
    )
    return "word/" + relationship.get("Target", "").lstrip("/")


def _assert_branded(path: Path) -> dict[str, object]:
    with zipfile.ZipFile(path) as package:
        names = set(package.namelist())
        document = _xml(package.read("word/document.xml"))
        rels = _xml(package.read("word/_rels/document.xml.rels"))
        section = document.find("w:body/w:sectPr", namespaces=NS)
        if section is None:
            raise RuntimeError(f"{path.name}: propriedades de seção ausentes.")
        page = section.find("w:pgSz", namespaces=NS)
        margins = section.find("w:pgMar", namespaces=NS)
        if page is None or margins is None:
            raise RuntimeError(f"{path.name}: página ou margens ausentes.")
        actual_page = (page.get(f"{W}w"), page.get(f"{W}h"))
        if actual_page != ("11906", "16838"):
            raise RuntimeError(f"{path.name}: não está em A4 retrato: {actual_page}")
        expected_margins = {
            "top": "1701",
            "right": "1134",
            "bottom": "1560",
            "left": "1134",
            "gutter": "0",
        }
        actual_margins = {key: margins.get(f"{W}{key}") for key in expected_margins}
        if actual_margins != expected_margins:
            raise RuntimeError(f"{path.name}: margens divergentes: {actual_margins}")

        header_ref = section.find("w:headerReference", namespaces=NS)
        footer_ref = section.find("w:footerReference", namespaces=NS)
        if header_ref is None or footer_ref is None:
            raise RuntimeError(f"{path.name}: cabeçalho ou rodapé ausente.")
        header_name = _relationship_target(rels, header_ref.get(f"{R}id"))
        footer_name = _relationship_target(rels, footer_ref.get(f"{R}id"))
        header = _xml(package.read(header_name))
        footer = _xml(package.read(footer_name))
        if not header.findall(".//a:blip", namespaces=NS):
            raise RuntimeError(f"{path.name}: imagem do timbre ausente no cabeçalho.")
        instructions = " ".join(footer.xpath(".//w:instrText/text()", namespaces=NS))
        footer_text = "".join(footer.xpath(".//w:t/text()", namespaces=NS))
        if "PAGE" not in instructions or "NUMPAGES" not in instructions or "Página" not in footer_text:
            raise RuntimeError(f"{path.name}: paginação 'Página X de Y' ausente.")
        media = sorted(name for name in names if name.startswith("word/media/"))
        if len(media) != 1:
            raise RuntimeError(f"{path.name}: deveria conter somente a imagem do timbre: {media}")
        paperhead = Image.open(BytesIO(package.read(media[0]))).convert("RGBA")
        white = Image.new("RGBA", paperhead.size, "white")
        composited = Image.alpha_composite(white, paperhead).convert("RGB")
        ink_bbox = ImageChops.difference(composited, Image.new("RGB", paperhead.size, "white")).getbbox()
        if ink_bbox is None:
            raise RuntimeError(f"{path.name}: a imagem do timbre ficou visualmente vazia.")
        return {
            "path": str(path.resolve()),
            "sha256": _sha256(path),
            "page": "A4 portrait",
            "margins_cm": [3.0, 2.0, 2.75, 2.0],
            "background_media": media,
            "background_visible_ink_bbox": list(ink_bbox),
            "pagination": "Página X de Y",
        }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--exe", required=True, type=Path)
    parser.add_argument("--modelo", required=True, type=Path)
    parser.add_argument("--trabalho", required=True, type=Path)
    parser.add_argument("--relatorio", required=True, type=Path)
    args = parser.parse_args()

    args.trabalho.mkdir(parents=True, exist_ok=True)
    run_root = Path(tempfile.mkdtemp(prefix="regressao_timbrado_", dir=args.trabalho))
    word_input = run_root / "peticao_word_controle.docx"
    control_image = run_root / "imagem_controle.png"
    before = _create_word_control(word_input, control_image)
    word_input_hash = _sha256(word_input)

    markdown_input = run_root / "peticao_markdown_controle.md"
    repeated = "\n\n".join(
        f"Parágrafo {index}. Texto de teste para verificar o papel timbrado e a paginação do documento."
        for index in range(1, 75)
    )
    markdown_input.write_text(
        "## PETIÇÃO DE CONTROLE\n\n### DOS FATOS\n\n" + repeated + "\n",
        encoding="utf-8",
    )

    word_output = _run(
        args.exe.resolve(),
        word_input,
        run_root / "saida_word",
        run_root / "localapp_word",
    )
    if _sha256(word_input) != word_input_hash:
        raise RuntimeError("O arquivo Word de entrada foi sobrescrito.")
    word_after = _assert_word_preserved(before, word_output, args.modelo)

    markdown_output = _run(
        args.exe.resolve(),
        markdown_input,
        run_root / "saida_markdown",
        run_root / "localapp_markdown",
    )
    intermediates = list((run_root / "saida_markdown").rglob("entrada_intermediaria.docx"))
    if len(intermediates) != 1:
        raise RuntimeError(f"Arquivo intermediário Markdown não localizado: {intermediates}")

    report = {
        "status": "ok",
        "executable": str(args.exe.resolve()),
        "model": str(args.modelo.resolve()),
        "model_sha256": _sha256(args.modelo),
        "run_root": str(run_root.resolve()),
        "word_input_unchanged": True,
        "word_page_size_header_footer_and_image_preserved": True,
        "word_margins_match_official_model": True,
        "word_snapshot": word_after,
        "markdown_intermediate": _assert_branded(intermediates[0]),
        "markdown_final": _assert_branded(markdown_output),
    }
    args.relatorio.parent.mkdir(parents=True, exist_ok=True)
    args.relatorio.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
