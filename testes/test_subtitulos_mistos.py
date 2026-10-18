import json
import tempfile
import unittest
from pathlib import Path

from scripts.classificar_regras import classificar, classificar_um


TEXTO_SUBTITULO = (
    "Pedidos de pesquisa patrimonial não equivalem a movimentação processual útil. "
    "Inexistência de direito automático à suspensão por 90 dias"
)


class SubtitulosMistosTests(unittest.TestCase):
    def paragrafo(self, texto, **extras):
        return {"id": 1, "texto": texto, "posicao_relativa": 0.6, **extras}

    def test_subtitulo_da_imagem_e_reconhecido_sem_alterar_texto(self):
        texto = "III.2. " + TEXTO_SUBTITULO
        item = classificar_um(self.paragrafo(texto))
        self.assertEqual(item["estilo_destino"], "FB - Título Secundário")
        self.assertFalse(item["necessita_analise_semantica"])
        self.assertEqual(item["texto"], texto)

    def test_hierarquia_mista_e_decimal_admitem_subtitulos_longos(self):
        for prefixo in ("III.2.", "IV.1", "III.2.1.", "3.2.", "3.2.1.", "III.2 –"):
            with self.subTest(prefixo=prefixo):
                item = classificar_um(self.paragrafo(prefixo + " " + TEXTO_SUBTITULO))
                self.assertEqual(item["estilo_destino"], "FB - Título Secundário")

    def test_numeracao_romana_invalida_nao_e_aceita_como_hierarquia(self):
        for prefixo in ("IIII.2.", "CIVIL.2.", "III..2."):
            with self.subTest(prefixo=prefixo):
                item = classificar_um(self.paragrafo(prefixo + " " + TEXTO_SUBTITULO))
                self.assertNotEqual(item["estilo_destino"], "FB - Título Secundário")

    def test_numeracao_nao_promove_itens_corridos_a_subtitulos(self):
        for texto in (
            "III.2. A parte autora apresenta os elementos necessários à análise do pedido.",
            "1.2. a produção de todas as provas admitidas;",
        ):
            with self.subTest(texto=texto):
                self.assertNotEqual(classificar_um(self.paragrafo(texto))["estilo_destino"], "FB - Título Secundário")
        item = classificar_um(self.paragrafo("III.2. " + TEXTO_SUBTITULO, estilo_original="Quote"))
        self.assertEqual(item["estilo_destino"], "FB - Citação")

    def test_docx_aplica_nivel_secundario_sem_duplicar_prefixo(self):
        from docx import Document
        from docx.oxml.ns import qn
        from scripts.aplicar_estilos import aplicar
        from scripts.extrair_documento import extrair

        root = Path(__file__).resolve().parents[1]
        with tempfile.TemporaryDirectory() as temp:
            entrada, saida = Path(temp) / "entrada.docx", Path(temp) / "saida.docx"
            doc = Document()
            doc.add_paragraph("III. DOS FUNDAMENTOS")
            doc.add_paragraph("III.2. " + TEXTO_SUBTITULO)
            doc.add_paragraph("A parte autora apresenta os elementos necessários à análise do pedido.")
            doc.save(entrada)
            original = entrada.read_bytes()
            config = json.loads((root / "modelos/estilos.json").read_text(encoding="utf-8"))
            regras = classificar(extrair(entrada))
            aplicar(entrada, root / "modelos/modelo_oficial.docx", regras, config, saida)
            p = Document(saida).paragraphs[1]
            self.assertEqual(p.style.name, "FB - Título Secundário")
            self.assertEqual(p.text, TEXTO_SUBTITULO)
            numpr = p._p.pPr.find(qn("w:numPr"))
            self.assertIsNotNone(numpr)
            self.assertEqual(numpr.find(qn("w:ilvl")).get(qn("w:val")), "1")
            self.assertEqual(entrada.read_bytes(), original)
