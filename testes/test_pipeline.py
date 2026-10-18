import base64
from io import BytesIO
import re
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path
import xml.etree.ElementTree as ET

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from scripts.classificar_regras import classificar
from scripts.comum import carregar_json, salvar_json
from scripts.criar_modelo import criar_modelo
from scripts.extrair_documento import extrair
from scripts.pipeline import executar
from scripts.validar_documento import validar


ROOT = Path(__file__).resolve().parents[1]


def _workspace_teste(base: str | Path) -> Path:
    workspace = Path(base) / "workspace_teste"
    modelos = workspace / "modelos"
    modelos.mkdir(parents=True, exist_ok=True)
    shutil.copy2(ROOT / "modelos" / "modelo_oficial.docx", modelos / "modelo_oficial.docx")
    shutil.copy2(ROOT / "modelos" / "estilos.json", modelos / "estilos.json")
    return workspace


def _amostra(path: Path):
    doc = Document()
    doc.add_paragraph("EXCELENTÍSSIMO SENHOR DOUTOR JUIZ DE DIREITO")
    doc.add_paragraph("Processo nº 0001234-56.2025.8.26.0001")
    doc.add_paragraph("DOS FATOS")
    doc.add_paragraph("A parte autora apresenta os fatos relevantes para o correto julgamento da demanda, preservando integralmente o conteúdo original.")
    doc.add_paragraph("Ante o exposto, requer:")
    doc.add_paragraph("1. a procedência integral do pedido;")
    doc.add_paragraph("Termos em que, pede deferimento.")
    doc.add_paragraph("São Paulo, 12 de julho de 2026.")
    doc.add_paragraph("Advogada - OAB/SP 123456")
    doc.save(path)


class PipelineTests(unittest.TestCase):
    def test_local_e_data_do_protocolo_recebem_estilo_de_fechamento(self):
        with tempfile.TemporaryDirectory() as d:
            entrada = Path(d) / "fechamento.docx"
            doc = Document()
            doc.add_paragraph("Termos em que, pede deferimento.")
            doc.add_paragraph("Aracaju/SE, data do protocolo.")
            doc.add_paragraph("FÁBIO AUGUSTO MENDONÇA BARRETO")
            doc.add_paragraph("OAB/SE nº 12.563")
            doc.save(entrada)
            original = entrada.read_bytes()
            resultado = executar(entrada, skip_pdf=True, work_root=_workspace_teste(d))
            final = Document(resultado["docx"])
            linha = next(p for p in final.paragraphs if p.text == "Aracaju/SE, data do protocolo.")
            self.assertEqual(linha.style.name, "FB - Fechamento")
            self.assertEqual(linha.alignment if linha.alignment is not None else linha.style.paragraph_format.alignment, 1)
            self.assertEqual(entrada.read_bytes(), original)

    def test_pipeline_interrompe_sem_artefatos_oficiais(self):
        with tempfile.TemporaryDirectory() as d:
            base = Path(d)
            entrada = base / "entrada.docx"
            _amostra(entrada)

            casos = (
                ("modelo_oficial.docx", ROOT / "modelos" / "estilos.json"),
                ("estilos.json", ROOT / "modelos" / "modelo_oficial.docx"),
            )
            for ausente, presente in casos:
                with self.subTest(ausente=ausente):
                    workspace = base / ausente.replace(".", "_")
                    modelos = workspace / "modelos"
                    modelos.mkdir(parents=True)
                    shutil.copy2(presente, modelos / presente.name)

                    with self.assertRaisesRegex(FileNotFoundError, "execução foi interrompida"):
                        executar(entrada, skip_pdf=True, work_root=workspace)
                    self.assertFalse((workspace / "saida" / "peticao_formatada.docx").exists())

    def test_pipeline_interrompe_com_regras_de_outro_escopo(self):
        with tempfile.TemporaryDirectory() as d:
            entrada = Path(d) / "entrada.docx"
            _amostra(entrada)
            workspace = _workspace_teste(d)
            config_path = workspace / "modelos" / "estilos.json"
            config = carregar_json(config_path)
            config["tipo_peca_suportado"] = "recurso"
            salvar_json(config_path, config)

            with self.assertRaisesRegex(ValueError, "escopo atual de petição inicial"):
                executar(entrada, skip_pdf=True, work_root=workspace)
            self.assertFalse((workspace / "saida" / "peticao_formatada.docx").exists())

    def test_extracao_e_regras(self):
        with tempfile.TemporaryDirectory() as d:
            entrada = Path(d) / "entrada.docx"; _amostra(entrada)
            estrutura = extrair(entrada); resultado = classificar(estrutura)
            self.assertEqual(len(estrutura["paragrafos"]), 9)
            self.assertEqual(resultado["classificacoes"][0]["estilo_destino"], "FB - Endereçamento")
            self.assertEqual(resultado["classificacoes"][1]["estilo_destino"], "FB - Identificação do Processo")

    def test_extracao_nao_infere_outro_tipo_por_palavras_do_documento(self):
        with tempfile.TemporaryDirectory() as d:
            entrada = Path(d) / "inicial_com_precedente.docx"
            doc = Document()
            doc.add_paragraph("AO JUÍZO DA VARA CÍVEL DA COMARCA DE ARACAJU - SE")
            doc.add_paragraph(
                "A petição inicial menciona precedente julgado em recurso de revista, "
                "sem que isso altere a natureza desta demanda."
            )
            doc.save(entrada)

            self.assertEqual(extrair(entrada)["tipo_peca"], "peticao_inicial")

    def test_pipeline_nao_bloqueia_cumprimento_de_sentenca(self):
        with tempfile.TemporaryDirectory() as d:
            entrada = Path(d) / "cumprimento.docx"
            doc = Document()
            doc.add_paragraph("AO JUÍZO DA 1ª VARA FEDERAL DE CURITIBA / PR")
            titulo = doc.add_paragraph("CUMPRIMENTO DE SENTENÇA")
            titulo.alignment = 1
            titulo.runs[0].bold = True
            doc.add_paragraph("Texto da fase processual posterior à petição inicial.")
            doc.save(entrada)
            original = entrada.read_bytes()
            workspace = _workspace_teste(d)

            self.assertEqual(extrair(entrada)["tipo_peca"], "peticao_inicial")
            resultado = executar(entrada, skip_pdf=True, work_root=workspace)
            self.assertTrue(Path(resultado["docx"]).exists())
            self.assertEqual(entrada.read_bytes(), original)

    def test_regras_priorizam_processo_qualificacao_e_assinatura(self):
        estrutura = {"documento": "cumprimento.docx", "paragrafos": [
            {"id": 1, "texto": "PROCESSO N.º 5036857-83.2025.4.04.7000", "hash": "a", "posicao_relativa": 0.03, "recuo_esquerdo_cm": 0, "negrito": True, "estilo_original": "Normal", "alinhamento": "LEFT (0)"},
            {"id": 2, "texto": "GABRIEL MACHADO DOS SANTOS, maior e capaz, médico, de CPF n.º 085.112.429-10, residente e domiciliado em Curitiba/PR, representado por Fábio Augusto Mendonça Barreto, advogado.", "hash": "b", "posicao_relativa": 0.18, "recuo_esquerdo_cm": 3, "negrito": False, "estilo_original": "Quote", "alinhamento": "JUSTIFY (3)"},
            {"id": 3, "texto": "Fábio Augusto Mendonça Barreto", "hash": "c", "posicao_relativa": 0.94, "recuo_esquerdo_cm": 0, "negrito": True, "estilo_original": "Normal", "alinhamento": "CENTER (1)"},
            {"id": 4, "texto": "OAB/SE nº 12.563 – OAB/SP nº 540.476", "hash": "d", "posicao_relativa": 0.97, "recuo_esquerdo_cm": 0, "negrito": True, "estilo_original": "Normal", "alinhamento": "CENTER (1)"},
        ]}
        resultado = classificar(estrutura)
        estilos = [item["estilo_destino"] for item in resultado["classificacoes"]]
        self.assertEqual(estilos[0], "FB - Identificação do Processo")
        self.assertEqual(estilos[1], "FB - Qualificação das Partes")
        self.assertEqual(estilos[2], "FB - Assinatura")
        self.assertEqual(estilos[3], "FB - Assinatura")
        self.assertFalse(resultado["ambiguos"])

    def test_regras_reconhecem_variantes_do_numero_da_oab(self):
        registros = (
            "OAB/SE nº 12.563",
            "OAB/SE n.º 12.563",
            "OAB/SE n° 12.563",
            "OAB/SE 12.563",
        )
        for indice, registro in enumerate(registros, 1):
            with self.subTest(registro=registro):
                estrutura = {"documento": "assinatura.docx", "paragrafos": [
                    {"id": 1, "texto": "Fábio Augusto Mendonça Barreto", "hash": f"nome-{indice}", "posicao_relativa": 0.94, "recuo_esquerdo_cm": 0, "negrito": True, "estilo_original": "Normal", "alinhamento": "CENTER (1)"},
                    {"id": 2, "texto": registro, "hash": f"oab-{indice}", "posicao_relativa": 0.97, "recuo_esquerdo_cm": 0, "negrito": True, "estilo_original": "Normal", "alinhamento": "CENTER (1)"},
                ]}
                resultado = classificar(estrutura)
                estilos = [item["estilo_destino"] for item in resultado["classificacoes"]]
                self.assertEqual(estilos, ["FB - Assinatura", "FB - Assinatura"])
                self.assertFalse(resultado["ambiguos"])

    def test_negrito_isolado_nao_produz_alta_confianca_de_titulo(self):
        estrutura = {"documento": "amostra.docx", "paragrafos": [{
            "id": 1,
            "texto": "Expressão destacada no corpo",
            "hash": "destaque",
            "posicao_relativa": 0.50,
            "recuo_esquerdo_cm": 0,
            "negrito": True,
            "estilo_original": "Normal",
            "alinhamento": "JUSTIFY (3)",
        }]}
        resultado = classificar(estrutura)
        item = resultado["classificacoes"][0]
        self.assertLess(item["confianca"], 0.80)
        self.assertTrue(item["necessita_analise_semantica"])

    def test_revisao_nao_substitui_regra_semantica_segura(self):
        with tempfile.TemporaryDirectory() as d:
            entrada = Path(d) / "entrada.docx"
            _amostra(entrada)
            estrutura = extrair(entrada)
            processo = estrutura["paragrafos"][1]
            respostas = Path(d) / "respostas.json"
            salvar_json(respostas, {"classificacoes": [{
                "id": processo["id"],
                "hash": processo["hash"],
                "estilo_destino": "FB - Título Principal",
                "confianca": 1.0,
                "justificativa": "Confirmação em lote incorreta.",
            }]})

            resultado = executar(entrada, semantica=respostas, skip_pdf=True, work_root=_workspace_teste(d))
            final = Document(resultado["docx"])
            processo_final = next(p for p in final.paragraphs if "Processo" in p.text)
            self.assertEqual(processo_final.style.name, "FB - Identificação do Processo")

    def test_auditoria_reprova_contradicoes_semanticas(self):
        with tempfile.TemporaryDirectory() as d:
            modelo = Path(d) / "modelo.docx"
            criar_modelo(modelo)
            doc = Document(modelo)
            corpo = doc._element.body
            for elemento in list(corpo):
                if elemento.tag != qn("w:sectPr"):
                    corpo.remove(elemento)
            doc.add_paragraph("PROCESSO N.º 5036857-83.2025.4.04.7000", style="FB - Título Principal")
            doc.add_paragraph(
                "GABRIEL MACHADO DOS SANTOS, maior e capaz, de CPF n.º 085.112.429-10, residente e domiciliado em Curitiba/PR, representado por seu advogado.",
                style="FB - Citação",
            )
            for numero in range(6):
                doc.add_paragraph(f"Parágrafo jurídico intermediário {numero}.", style="FB - Corpo")
            doc.add_paragraph("Fábio Augusto Mendonça Barreto", style="FB - Título Secundário")
            doc.add_paragraph("OAB/SE nº 12.563", style="FB - Título Secundário")
            saida = Path(d) / "incorreto.docx"
            doc.save(saida)

            auditoria = validar(saida, carregar_json(ROOT / "modelos" / "estilos.json"))
            codigos = {item["codigo"] for item in auditoria["problemas"]}
            self.assertIn("identificacao_processo_com_estilo_incorreto", codigos)
            self.assertIn("qualificacao_com_estilo_citacao", codigos)
            self.assertIn("nome_assinatura_com_estilo_incorreto", codigos)
            self.assertIn("assinatura_com_estilo_incorreto", codigos)
            self.assertFalse(auditoria["resumo"]["aprovado"])

    def test_resumo_inicial(self):
        estrutura = {"documento": "amostra.docx", "paragrafos": [
            {"id": 1, "texto": "RESUMO: DIREITO À MORADIA.", "hash": "x", "posicao_relativa": 0.01, "recuo_esquerdo_cm": 0, "negrito": False},
            {"id": 2, "texto": "- Tema 325 da TNU.", "hash": "y", "posicao_relativa": 0.02, "recuo_esquerdo_cm": 0, "negrito": False},
            {"id": 3, "texto": "- Legislação relevante.", "hash": "z", "posicao_relativa": 0.03, "recuo_esquerdo_cm": 0, "negrito": False},
        ]}
        resultado = classificar(estrutura)
        self.assertTrue(all(x["estilo_destino"] == "FB - Resumo" for x in resultado["classificacoes"]))

    def test_estilos_de_resumo_no_modelo(self):
        with tempfile.TemporaryDirectory() as d:
            modelo = Path(d) / "modelo.docx"
            criar_modelo(modelo)
            doc = Document(modelo)
            resumo = doc.styles["FB - Resumo"]
            self.assertEqual(resumo.font.name, "Segoe UI Variable Text")
            self.assertEqual(resumo.font.size.pt, 10)
            self.assertAlmostEqual(resumo.paragraph_format.left_indent.cm, 9, places=2)
            acao = doc.styles["FB - Identificação da Ação"]
            self.assertAlmostEqual(acao.paragraph_format.space_before.pt, 12, places=2)
            processo = doc.styles["FB - Identificação do Processo"]
            self.assertTrue(processo.font.bold)
            self.assertIsNone(processo.element.pPr.find(qn("w:pBdr")))
            tabela = doc.styles["FB - Tabela"]
            self.assertEqual(tabela.paragraph_format.first_line_indent.cm, 0)
            self.assertEqual(tabela.paragraph_format.space_before.pt, 0)
            self.assertEqual(tabela.paragraph_format.space_after.pt, 0)
            for nome in ("FB - Citação", "FB - Lista", "FB - Pedido"):
                estilo = doc.styles[nome]
                self.assertAlmostEqual(estilo.paragraph_format.left_indent.cm, 3, places=2)
            self.assertEqual(doc.styles["FB - Lista"].paragraph_format.first_line_indent.cm, 0)
            self.assertEqual(doc.styles["FB - Pedido"].paragraph_format.first_line_indent.cm, 0)
            principal = doc.styles["FB - Título Principal"]
            secundario = doc.styles["FB - Título Secundário"]
            self.assertTrue(principal.font.bold)
            self.assertTrue(secundario.font.bold)
            self.assertIsNotNone(principal.element.pPr.find(qn("w:numPr")))
            self.assertIsNotNone(secundario.element.pPr.find(qn("w:numPr")))
            self.assertAlmostEqual(secundario.paragraph_format.left_indent.cm, 1, places=2)
            self.assertAlmostEqual(secundario.paragraph_format.first_line_indent.cm, 1, places=2)

    def test_catalogo_semantico_cobre_todos_os_estilos(self):
        config = carregar_json(Path(__file__).parents[1] / "modelos" / "estilos.json")
        catalogo = config["catalogo_semantico"]
        formatacao = config["formatacao_estilos"]
        catalogo_md = (Path(__file__).parents[1] / "modelos" / "catalogo_estilos.md").read_text(encoding="utf-8")
        self.assertEqual(set(catalogo), set(config["estilos_permitidos"]))
        self.assertEqual(set(formatacao), set(config["estilos_permitidos"]))
        self.assertEqual(config["tipo_peca_suportado"], "peticao_inicial")
        self.assertEqual(
            set(config["regras_contextuais"]),
            {"peticao_inicial", "normalizacao_estrutural"},
        )
        self.assertTrue(all(f"`{estilo}`" in catalogo_md for estilo in config["estilos_permitidos"]))
        for estilo in config["estilos_permitidos"]:
            entrada = catalogo[estilo]
            self.assertTrue(entrada["descricao"])
            self.assertTrue(entrada["aplicar_quando"])
            self.assertTrue(entrada["nao_aplicar_quando"])
            self.assertTrue(entrada["exemplos"])
        self.assertTrue(any("pedido" in item.lower() for item in catalogo["FB - Corpo"]["nao_aplicar_quando"]))
        self.assertTrue(any("bloco DOS PEDIDOS" in item for item in catalogo["FB - Pedido"]["aplicar_quando"]))

    def test_formatacao_catalogada_reflete_modelo_oficial(self):
        root = Path(__file__).parents[1]
        config = carregar_json(root / "modelos" / "estilos.json")
        doc = Document(root / "modelos" / "modelo_oficial.docx")
        alinhamentos = {"LEFT (0)": "esquerda", "JUSTIFY (3)": "justificado", "CENTER (1)": "centralizado"}

        def alinhamento_efetivo(estilo):
            atual = estilo
            while atual is not None:
                if atual.paragraph_format.alignment is not None:
                    return atual.paragraph_format.alignment
                atual = atual.base_style
            return "LEFT (0)"  # alinhamento padrão do Word: esquerda

        for nome, esperado in config["formatacao_estilos"].items():
            estilo = doc.styles[nome]
            pf = estilo.paragraph_format
            atual = {
                "fonte": estilo.font.name,
                "tamanho_pt": estilo.font.size.pt if estilo.font.size else None,
                "negrito": bool(estilo.font.bold),
                "italico": bool(estilo.font.italic),
                "alinhamento": alinhamentos.get(str(alinhamento_efetivo(estilo)), str(alinhamento_efetivo(estilo))),
                "recuo_esquerdo_cm": pf.left_indent.cm if pf.left_indent else 0,
                "recuo_direito_cm": pf.right_indent.cm if pf.right_indent else 0,
                "recuo_primeira_linha_cm": pf.first_line_indent.cm if pf.first_line_indent else 0,
                "espaco_antes_pt": pf.space_before.pt if pf.space_before else 0,
                "espaco_depois_pt": pf.space_after.pt if pf.space_after else 0,
                "entrelinhas": pf.line_spacing,
                "manter_com_proximo": bool(pf.keep_with_next),
            }
            for chave, valor in atual.items():
                esperado_valor = esperado.get(chave, 0)
                if isinstance(valor, float) or isinstance(esperado_valor, float):
                    self.assertAlmostEqual(float(valor), float(esperado_valor), places=2, msg=f"{nome}: {chave}")
                else:
                    self.assertEqual(valor, esperado_valor, f"{nome}: {chave}")

    def test_modelo_novo_exibe_apenas_estilos_autoriais(self):
        root = Path(__file__).parents[1]
        config = carregar_json(root / "modelos" / "estilos.json")
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

        with tempfile.TemporaryDirectory() as d:
            destino = Path(d) / "modelo.docx"
            criar_modelo(destino)
            with zipfile.ZipFile(destino) as pacote:
                styles = ET.fromstring(pacote.read("word/styles.xml"))
                visible = []
                for style in styles.findall("w:style", ns):
                    if style.find("w:semiHidden", ns) is None:
                        name = style.find("w:name", ns)
                        if name is not None:
                            visible.append(name.get("{%s}val" % ns["w"]))

                settings = ET.fromstring(pacote.read("word/settings.xml"))
                pane_filter = settings.find("w:stylePaneFormatFilter", ns)

            self.assertEqual(visible, config["estilos_permitidos"])
            self.assertIsNotNone(pane_filter)
            self.assertEqual(pane_filter.get("{%s}customStyles" % ns["w"]), "1")

            # Um prefixo citado em mc:Ignorable precisa continuar declarado.
            # Caso contrário, o Word acusa o DOCX como aparentemente corrompido.
            with zipfile.ZipFile(destino) as pacote:
                for parte in ("word/styles.xml", "word/settings.xml"):
                    xml = pacote.read(parte).decode("utf-8")
                    ignoraveis = re.search(r'(?:mc|ns\d+):Ignorable="([^"]+)"', xml)
                    if ignoraveis:
                        for prefixo in ignoraveis.group(1).split():
                            self.assertRegex(xml, rf'xmlns:{re.escape(prefixo)}="[^"]+"')

    def test_saida_final_exibe_apenas_estilos_autoriais(self):
        root = Path(__file__).parents[1]
        config = carregar_json(root / "modelos" / "estilos.json")
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

        with tempfile.TemporaryDirectory() as d:
            entrada = Path(d) / "original.docx"
            _amostra(entrada)
            resultado = executar(entrada, skip_pdf=True, work_root=_workspace_teste(d))

            with zipfile.ZipFile(resultado["docx"]) as pacote:
                styles = ET.fromstring(pacote.read("word/styles.xml"))
                visible = []
                style_order = {
                    name: index
                    for index, name in enumerate(
                        (
                            "name", "aliases", "basedOn", "next", "link", "autoRedefine",
                            "hidden", "uiPriority", "semiHidden", "unhideWhenUsed", "qFormat",
                            "locked", "personal", "personalCompose", "personalReply", "rsid",
                            "pPr", "rPr", "tblPr", "trPr", "tcPr", "tblStylePr",
                        )
                    )
                }
                for style in styles.findall("w:style", ns):
                    child_names = [child.tag.rsplit("}", 1)[-1] for child in list(style)]
                    self.assertEqual(
                        [style_order[name] for name in child_names],
                        sorted(style_order[name] for name in child_names),
                        f"ordem XML inválida no estilo {style.get('{%s}type' % ns['w'])}",
                    )
                    if style.find("w:semiHidden", ns) is None:
                        name = style.find("w:name", ns)
                        if name is not None:
                            visible.append(name.get("{%s}val" % ns["w"]))

                settings = ET.fromstring(pacote.read("word/settings.xml"))
                pane_filter = settings.find("w:stylePaneFormatFilter", ns)

            self.assertEqual(visible, config["estilos_permitidos"])
            self.assertIsNotNone(pane_filter)
            self.assertEqual(pane_filter.get("{%s}customStyles" % ns["w"]), "1")
            settings_names = [child.tag.rsplit("}", 1)[-1] for child in list(settings)]
            self.assertLess(settings_names.index("stylePaneFormatFilter"), settings_names.index("defaultTabStop"))

    def test_saida_usa_fonte_oficial_nos_estilos_e_no_corpo(self):
        root = Path(__file__).parents[1]
        config = carregar_json(root / "modelos" / "estilos.json")
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

        with tempfile.TemporaryDirectory() as d:
            entrada = Path(d) / "original.docx"
            _amostra(entrada)
            resultado = executar(entrada, skip_pdf=True, work_root=_workspace_teste(d))
            final = Document(resultado["docx"])

            for nome in config["estilos_permitidos"]:
                self.assertEqual(final.styles[nome].font.name, "Segoe UI Variable Text", nome)

            with zipfile.ZipFile(resultado["docx"]) as pacote:
                document = ET.fromstring(pacote.read("word/document.xml"))
                for fonts in document.findall(".//w:rFonts", ns):
                    valores = [fonts.get("{%s}%s" % (ns["w"], atributo)) for atributo in ("ascii", "hAnsi", "eastAsia", "cs")]
                    self.assertTrue(all(not valor or valor == "Segoe UI Variable Text" for valor in valores), valores)

    def test_titulos_tem_linha_de_separacao_anterior(self):
        doc = Document(Path(__file__).parents[1] / "modelos" / "modelo_oficial.docx")
        self.assertAlmostEqual(doc.styles["FB - Identificação da Ação"].paragraph_format.space_before.pt, 12, places=2)
        for nome in ("FB - Título Principal", "FB - Título Secundário"):
            self.assertAlmostEqual(doc.styles[nome].paragraph_format.space_before.pt, 36, places=2)

    def test_bloco_final_nao_tem_espacamento(self):
        doc = Document(Path(__file__).parents[1] / "modelos" / "modelo_oficial.docx")
        for nome in ("FB - Fechamento", "FB - Assinatura"):
            pf = doc.styles[nome].paragraph_format
            self.assertAlmostEqual(pf.space_before.pt if pf.space_before else 0, 0, places=2)
            self.assertAlmostEqual(pf.space_after.pt if pf.space_after else 0, 0, places=2)

    def test_titulos_usam_uma_lista_decimal_multinivel_e_ficam_em_negrito(self):
        with tempfile.TemporaryDirectory() as d:
            entrada = Path(d) / "original.docx"
            _amostra(entrada)
            doc = Document(entrada)
            principal_origem = doc.paragraphs[0]
            principal_origem.style = doc.styles["Heading 1"]
            principal_origem.text = "2. DOS FATOS"
            principal_origem.runs[0].bold = False
            secundario_origem = doc.paragraphs[2]
            secundario_origem.style = doc.styles["Heading 2"]
            secundario_origem.text = "2.1. DA PROVA DOCUMENTAL"
            secundario_origem.runs[0].bold = False
            numbering = doc.part.numbering_part.element
            abstract_id = str(max([int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))] + [0]) + 1)
            num_id = str(max([int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))] + [0]) + 1)
            abstract = OxmlElement("w:abstractNum"); abstract.set(qn("w:abstractNumId"), abstract_id)
            lvl = OxmlElement("w:lvl"); lvl.set(qn("w:ilvl"), "0")
            for tag, value in (("start", "1"), ("numFmt", "upperRoman"), ("lvlText", "%1."), ("lvlJc", "left")):
                node = OxmlElement(f"w:{tag}"); node.set(qn("w:val"), value); lvl.append(node)
            lvl_ppr = OxmlElement("w:pPr"); lvl_ind = OxmlElement("w:ind"); lvl_ind.set(qn("w:left"), "720"); lvl_ind.set(qn("w:hanging"), "360"); lvl_ppr.append(lvl_ind); lvl.append(lvl_ppr)
            abstract.append(lvl); numbering.append(abstract)
            num = OxmlElement("w:num"); num.set(qn("w:numId"), num_id); abstract_ref = OxmlElement("w:abstractNumId"); abstract_ref.set(qn("w:val"), abstract_id); num.append(abstract_ref); numbering.append(num)
            ppr = principal_origem._p.get_or_add_pPr(); numpr = OxmlElement("w:numPr"); ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "0"); numid = OxmlElement("w:numId"); numid.set(qn("w:val"), num_id); numpr.extend([ilvl, numid]); ppr.append(numpr)
            doc.save(entrada)
            resultado = executar(entrada, skip_pdf=True, work_root=_workspace_teste(d))
            self.assertTrue(resultado["aprovado"])
            self.assertEqual(resultado["ambiguos"], 0)
            final = Document(resultado["docx"])
            titulo = next(p for p in final.paragraphs if p.text == "DOS FATOS")
            subtitulo = next(p for p in final.paragraphs if p.text == "DA PROVA DOCUMENTAL")
            self.assertEqual(titulo.style.name, "FB - Título Principal")
            self.assertEqual(subtitulo.style.name, "FB - Título Secundário")
            self.assertTrue(titulo.style.font.bold)
            self.assertTrue(subtitulo.style.font.bold)
            self.assertTrue(all(run.bold for run in titulo.runs + subtitulo.runs))

            numpr_principal = titulo._p.pPr.find(qn("w:numPr"))
            numpr_secundario = subtitulo._p.pPr.find(qn("w:numPr"))
            self.assertIsNotNone(numpr_principal)
            self.assertIsNotNone(numpr_secundario)
            numid = numpr_principal.find(qn("w:numId")).get(qn("w:val"))
            self.assertEqual(numid, numpr_secundario.find(qn("w:numId")).get(qn("w:val")))
            self.assertEqual(numpr_principal.find(qn("w:ilvl")).get(qn("w:val")), "0")
            self.assertEqual(numpr_secundario.find(qn("w:ilvl")).get(qn("w:val")), "1")
            for estilo, nivel in ((titulo.style, "0"), (subtitulo.style, "1")):
                numpr_estilo = estilo.element.pPr.find(qn("w:numPr"))
                self.assertEqual(numpr_estilo.find(qn("w:numId")).get(qn("w:val")), numid)
                self.assertEqual(numpr_estilo.find(qn("w:ilvl")).get(qn("w:val")), nivel)
            numbering = final.part.numbering_part.element
            num = next(x for x in numbering.findall(qn("w:num")) if x.get(qn("w:numId")) == numid)
            abstract_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
            abstract = next(x for x in numbering.findall(qn("w:abstractNum")) if x.get(qn("w:abstractNumId")) == abstract_id)
            self.assertEqual(len(abstract.findall(qn("w:lvl"))), 9)
            nivel_principal = next(x for x in abstract.findall(qn("w:lvl")) if x.get(qn("w:ilvl")) == "0")
            nivel_secundario = next(x for x in abstract.findall(qn("w:lvl")) if x.get(qn("w:ilvl")) == "1")
            self.assertEqual(nivel_principal.find(qn("w:numFmt")).get(qn("w:val")), "decimal")
            self.assertEqual(nivel_principal.find(qn("w:lvlText")).get(qn("w:val")), "%1")
            self.assertEqual(nivel_secundario.find(qn("w:numFmt")).get(qn("w:val")), "decimal")
            self.assertEqual(nivel_secundario.find(qn("w:lvlText")).get(qn("w:val")), "%1.%2")
            for nivel in (nivel_principal, nivel_secundario):
                rpr = nivel.find(qn("w:rPr"))
                self.assertEqual(rpr.find(qn("w:sz")).get(qn("w:val")), "24")
                self.assertIsNotNone(rpr.find(qn("w:b")))
            ind_principal = nivel_principal.find(qn("w:pPr")).find(qn("w:ind"))
            ind_secundario = nivel_secundario.find(qn("w:pPr")).find(qn("w:ind"))
            self.assertEqual((ind_principal.get(qn("w:left")), ind_principal.get(qn("w:hanging"))), ("283", "283"))
            self.assertEqual((ind_secundario.get(qn("w:left")), ind_secundario.get(qn("w:firstLine"))), ("567", "567"))
            self.assertIsNone(ind_secundario.get(qn("w:hanging")))

    def test_solicitacao_semantica_inclui_catalogo(self):
        with tempfile.TemporaryDirectory() as d:
            entrada = Path(d) / "original.docx"; _amostra(entrada)
            workspace = _workspace_teste(d)
            resultado = executar(entrada, skip_pdf=True, work_root=workspace)
            solicitacoes = carregar_json(workspace / "trabalho" / "solicitacoes_semanticas.json")
            self.assertEqual(set(solicitacoes["catalogo_semantico"]), set(solicitacoes["estilos_permitidos"]))
            self.assertEqual(set(solicitacoes["formatacao_estilos"]), set(solicitacoes["estilos_permitidos"]))
            self.assertIn("catalogo_semantico", solicitacoes["instrucoes"])
            self.assertIn("tipo_peca", solicitacoes)
            self.assertIn("regras_contextuais", solicitacoes)
            self.assertTrue(Path(resultado["relatorio"]).exists())

    def test_revisao_semantica_resolve_pendencias_e_aprova_auditoria(self):
        with tempfile.TemporaryDirectory() as d:
            entrada = Path(d) / "original.docx"
            _amostra(entrada)
            workspace = _workspace_teste(d)
            primeiro = executar(entrada, skip_pdf=True, work_root=workspace)
            solicitacoes = carregar_json(workspace / "trabalho" / "solicitacoes_semanticas.json")
            self.assertGreater(primeiro["ambiguos"], 0)
            self.assertFalse(primeiro["aprovado"])

            respostas = workspace / "trabalho" / "respostas_semanticas.json"
            salvar_json(respostas, {"classificacoes": [
                {
                    "id": item["id"],
                    "hash": item["hash"],
                    "estilo_destino": item["pre_classificacao"],
                    "confianca": 1.0,
                    "justificativa": "Confirmado no teste de revisão.",
                }
                for item in solicitacoes["itens"]
            ]})

            revisado = executar(entrada, semantica=respostas, skip_pdf=True, work_root=workspace)
            self.assertEqual(revisado["ambiguos"], 0)
            self.assertTrue(revisado["aprovado"])
            novas_solicitacoes = carregar_json(workspace / "trabalho" / "solicitacoes_semanticas.json")
            self.assertEqual(novas_solicitacoes["itens"], [])

    def test_revisao_semantica_rejeita_resposta_de_texto_alterado(self):
        with tempfile.TemporaryDirectory() as d:
            entrada = Path(d) / "original.docx"
            _amostra(entrada)
            workspace = _workspace_teste(d)
            primeiro = executar(entrada, skip_pdf=True, work_root=workspace)
            solicitacoes = carregar_json(workspace / "trabalho" / "solicitacoes_semanticas.json")
            item = solicitacoes["itens"][0]
            respostas = workspace / "trabalho" / "respostas_semanticas.json"
            salvar_json(respostas, {"classificacoes": [{
                "id": item["id"],
                "hash": "hash-de-versao-antiga",
                "estilo_destino": item["pre_classificacao"],
                "confianca": 1.0,
                "justificativa": "Resposta obsoleta.",
            }]})

            revisado = executar(entrada, semantica=respostas, skip_pdf=True, work_root=workspace)
            self.assertEqual(revisado["ambiguos"], primeiro["ambiguos"])
            self.assertFalse(revisado["aprovado"])

    def test_classificador_ignora_tipo_antigo_e_mantem_escopo_inicial(self):
        estrutura = {"documento": "inicial.docx", "tipo_peca": "recurso_de_revista", "paragrafos": [
            {"id": 1, "texto": "Recorrente: MARIA DA SILVA", "hash": "a", "posicao_relativa": 0.05, "recuo_esquerdo_cm": 0, "negrito": False},
            {"id": 2, "texto": "AO VICE-PRESIDENTE DO TRIBUNAL REGIONAL DO TRABALHO", "hash": "b", "posicao_relativa": 0.06, "recuo_esquerdo_cm": 0, "negrito": False},
        ]}
        resultado = classificar(estrutura)
        self.assertEqual(resultado["tipo_peca"], "peticao_inicial")
        self.assertNotEqual(resultado["classificacoes"][0]["estilo_destino"], "FB - Identificação do Processo")
        self.assertNotEqual(resultado["classificacoes"][1]["estilo_destino"], "FB - Endereçamento")

    def test_peticao_inicial_tem_normalizacao_estrutural(self):
        with tempfile.TemporaryDirectory() as d:
            entrada = Path(d) / "peticao_inicial.docx"
            doc = Document()
            doc.add_paragraph("AO JUÍZO DA VARA CÍVEL DA COMARCA DE ARACAJU - SE")
            doc.add_paragraph("Processo nº 0001234-56.2025.8.25.0001")
            doc.add_paragraph("FÁBIO DOS SANTOS, já qualificado, propõe a presente AÇÃO DE OBRIGAÇÃO DE FAZER.")
            doc.add_paragraph("## 2. FUNDAMENTOS DO PEDIDO")
            doc.add_paragraph("### 2.1. DA TUTELA DE URGÊNCIA")
            doc.add_paragraph("> **Trecho literal da decisão mencionada.**")
            doc.add_paragraph("| Critério | Resultado |")
            doc.add_paragraph("| --- | --- |")
            doc.add_paragraph("| Probabilidade do direito | Demonstrada |")
            doc.add_paragraph("## 8. PEDIDOS")
            doc.add_paragraph("Diante do exposto, requer-se:")
            doc.add_paragraph("1. a concessão da tutela de urgência;")
            doc.add_paragraph("Aracaju/SE, 22 de julho de 2026.")
            doc.add_paragraph("")
            assinatura = doc.add_table(rows=1, cols=1)
            assinatura.cell(0, 0).text = "Fábio Augusto Mendonça Barreto\nOAB/SE nº 12.563"
            doc.save(entrada)

            estrutura = extrair(entrada)
            self.assertEqual(estrutura["tipo_peca"], "peticao_inicial")
            regras = classificar(estrutura)["classificacoes"]
            self.assertEqual(regras[0]["estilo_destino"], "FB - Endereçamento")
            self.assertEqual(regras[2]["estilo_destino"], "FB - Qualificação das Partes")

            resultado = executar(entrada, skip_pdf=True, work_root=_workspace_teste(d))
            self.assertTrue(resultado["aprovado"])
            self.assertEqual(resultado["ambiguos"], 0)
            final = Document(resultado["docx"])
            local_data = next(p for p in final.paragraphs if p.text == "Aracaju/SE, 22 de julho de 2026.")
            self.assertEqual(local_data.style.name, "FB - Fechamento")
            textos = [p.text for p in final.paragraphs]
            self.assertFalse(any(t.lstrip().startswith(("##", "###", ">")) or "**" in t for t in textos))
            self.assertEqual(len(final.tables), 2)
            titulo = next(p for p in final.paragraphs if p.text == "FUNDAMENTOS DO PEDIDO")
            self.assertEqual(titulo.style.name, "FB - Título Principal")
            self.assertTrue(all(run.bold for run in titulo.runs))
            self.assertEqual(titulo._p.pPr.find(qn("w:numPr")).find(qn("w:ilvl")).get(qn("w:val")), "0")
            pedido = next(p for p in final.paragraphs if "concessão da tutela" in p.text)
            self.assertEqual(pedido.style.name, "FB - Pedido")
            self.assertTrue(pedido.text.startswith("1."))
            assinatura_final = next(t for t in final.tables if "OAB/SE" in t.cell(0, 0).text)
            self.assertTrue(all(p.style.name == "FB - Assinatura" for p in assinatura_final.cell(0, 0).paragraphs))

    def test_no_caso_dos_municipios_nao_e_processo(self):
        estrutura = {"documento": "amostra.docx", "paragrafos": [
            {"id": 1, "texto": "No caso dos Municípios, aplica-se a regra legal pertinente.", "hash": "a", "posicao_relativa": 0.5, "recuo_esquerdo_cm": 0, "negrito": False},
            {"id": 2, "texto": "Processo nº 0001234-56.2025.8.26.0001", "hash": "b", "posicao_relativa": 0.05, "recuo_esquerdo_cm": 0, "negrito": False},
        ]}
        resultado = classificar(estrutura)["classificacoes"]
        self.assertNotEqual(resultado[0]["estilo_destino"], "FB - Identificação do Processo")
        self.assertEqual(resultado[1]["estilo_destino"], "FB - Identificação do Processo")

    def test_estilo_quote_e_citacao(self):
        estrutura = {"documento": "amostra.docx", "paragrafos": [
            {"id": 1, "texto": "Trecho jurisprudencial transcrito.", "hash": "a", "posicao_relativa": 0.5, "recuo_esquerdo_cm": 0, "negrito": False, "estilo_original": "Quote"},
        ]}
        self.assertEqual(classificar(estrutura)["classificacoes"][0]["estilo_destino"], "FB - Citação")

    def test_estilos_estruturais_de_titulo_prevalecem(self):
        estrutura = {"documento": "amostra.docx", "paragrafos": [
            {"id": 1, "texto": "DOS PEDIDOS E DOS REQUERIMENTOS", "hash": "a", "posicao_relativa": 0.9, "recuo_esquerdo_cm": 0, "negrito": True, "estilo_original": "Heading 1"},
            {"id": 2, "texto": "DA NORMA GERAL", "hash": "b", "posicao_relativa": 0.5, "recuo_esquerdo_cm": 0, "negrito": True, "estilo_original": "Heading 2"},
        ]}
        resultado = classificar(estrutura)["classificacoes"]
        self.assertEqual(resultado[0]["estilo_destino"], "FB - Título Principal")
        self.assertEqual(resultado[1]["estilo_destino"], "FB - Título Secundário")

    def test_introducao_e_itens_do_bloco_de_pedidos(self):
        estrutura = {"documento": "amostra.docx", "paragrafos": [
            {"id": 1, "texto": "DOS PEDIDOS E DOS REQUERIMENTOS", "hash": "a", "posicao_relativa": 0.8, "recuo_esquerdo_cm": 0, "negrito": True, "estilo_original": "Heading 1", "possui_numeracao": False},
            {"id": 2, "texto": "Diante do exposto, requer-se a V. Exa.:", "hash": "b", "posicao_relativa": 0.81, "recuo_esquerdo_cm": 0, "negrito": False, "estilo_original": "Normal", "possui_numeracao": False},
            {"id": 3, "texto": "o recebimento e processamento da petição;", "hash": "c", "posicao_relativa": 0.82, "recuo_esquerdo_cm": 3, "negrito": False, "estilo_original": "Normal", "possui_numeracao": True},
            {"id": 4, "texto": "a procedência do pedido;", "hash": "d", "posicao_relativa": 0.83, "recuo_esquerdo_cm": 3, "negrito": False, "estilo_original": "Normal", "possui_numeracao": True},
        ]}
        resultado = classificar(estrutura)["classificacoes"]
        self.assertEqual(resultado[1]["estilo_destino"], "FB - Corpo")
        self.assertEqual(resultado[2]["estilo_destino"], "FB - Pedido")
        self.assertEqual(resultado[3]["estilo_destino"], "FB - Pedido")

    def test_qualificacao_e_acao(self):
        estrutura = {"documento": "amostra.docx", "paragrafos": [
            {"id": 1, "texto": "AÇÃO PELO RITO DO JUIZADO ESPECIAL", "hash": "a", "posicao_relativa": 0.10, "recuo_esquerdo_cm": 0, "negrito": True},
            {"id": 2, "texto": "MARIA, já qualificada nos autos, vem respeitosamente à presença de Vossa Excelência.", "hash": "b", "posicao_relativa": 0.15, "recuo_esquerdo_cm": 0, "negrito": False},
        ]}
        resultado = classificar(estrutura)
        self.assertEqual(resultado["classificacoes"][0]["estilo_destino"], "FB - Identificação da Ação")
        self.assertEqual(resultado["classificacoes"][1]["estilo_destino"], "FB - Qualificação das Partes")

    def test_qualificacao_restringe_negrito_a_trechos_relevantes(self):
        with tempfile.TemporaryDirectory() as d:
            entrada = Path(d) / "qualificacao.docx"
            doc = Document()
            paragrafo = doc.add_paragraph()
            run = paragrafo.add_run(
                "JORGE LUIZ DE BARROS TORRESI, já qualificado, nos autos da ação que move contra "
                "a FAZENDA DO ESTADO DE SÃO PAULO, apresenta EMBARGOS DE DECLARAÇÃO em face da sentença."
            )
            run.bold = True
            doc.save(entrada)
            resultado = executar(entrada, skip_pdf=True, work_root=_workspace_teste(d))
            final = Document(resultado["docx"])
            p = final.paragraphs[0]
            negrito = "".join(run.text for run in p.runs if run.bold)
            self.assertIn("JORGE LUIZ DE BARROS TORRESI", negrito)
            self.assertIn("FAZENDA DO ESTADO DE SÃO PAULO", negrito)
            self.assertIn("EMBARGOS DE DECLARAÇÃO", negrito)
            self.assertIn("ação", negrito)
            self.assertIn("sentença", negrito)
            self.assertNotIn("já qualificado", negrito)
            self.assertNotIn("nos autos", negrito)

    def test_corpo_remove_negrito_direto(self):
        with tempfile.TemporaryDirectory() as d:
            entrada = Path(d) / "corpo.docx"
            doc = Document()
            p = doc.add_paragraph()
            p.add_run("Texto jurídico corrido que não deve permanecer em negrito porque desenvolve fatos, fundamentos, argumentos e transições dentro da peça processual, sem exercer função de título, subtítulo, pedido, assinatura ou qualquer outro papel estrutural.").bold = True
            doc.save(entrada)
            resultado = executar(entrada, skip_pdf=True, work_root=_workspace_teste(d))
            final = Document(resultado["docx"])
            self.assertTrue(all(not bool(run.bold) for run in final.paragraphs[0].runs))

    def test_imagem_de_assinatura_fica_centralizada(self):
        with tempfile.TemporaryDirectory() as d:
            entrada = Path(d) / "assinatura.docx"
            doc = Document()
            doc.add_paragraph("Termos em que, pede deferimento.")
            imagem = doc.add_paragraph()
            imagem.add_run().add_picture(BytesIO(base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII=")))
            doc.add_paragraph("Fábio Augusto Mendonça Barreto")
            doc.add_paragraph("OAB/SE nº 12.563")
            doc.save(entrada)
            resultado = executar(entrada, skip_pdf=True, work_root=_workspace_teste(d))
            final = Document(resultado["docx"])
            paragrafos_com_imagem = [p for p in final.paragraphs if p._p.xpath(".//w:drawing")]
            self.assertEqual(len(paragrafos_com_imagem), 1)
            self.assertEqual(paragrafos_com_imagem[0].style.name, "FB - Assinatura")
            assinatura = [p for p in final.paragraphs if p.text in {"Fábio Augusto Mendonça Barreto", "OAB/SE nº 12.563"}]
            self.assertEqual([p.style.name for p in assinatura], ["FB - Assinatura", "FB - Assinatura"])

    def test_pipeline_preserva_texto_e_original(self):
        with tempfile.TemporaryDirectory() as d:
            entrada = Path(d) / "original.docx"; _amostra(entrada)
            antes = entrada.read_bytes()
            resultado = executar(entrada, skip_pdf=True, work_root=_workspace_teste(d))
            self.assertEqual(resultado["tipo_peca"], "peticao_inicial")
            self.assertEqual(entrada.read_bytes(), antes)
            final = Document(resultado["docx"]); original = Document(entrada)
            self.assertEqual([p.text for p in final.paragraphs], [p.text for p in original.paragraphs])
            config = carregar_json(Path(__file__).parents[1] / "modelos" / "estilos.json")
            self.assertTrue(all(p.style.name in config["estilos_permitidos"] for p in final.paragraphs))
            self.assertTrue(Path(resultado["relatorio"]).exists())

    def test_identificacao_do_processo_fica_negrita_na_saida(self):
        with tempfile.TemporaryDirectory() as d:
            entrada = Path(d) / "original.docx"; _amostra(entrada)
            resultado = executar(entrada, skip_pdf=True, work_root=_workspace_teste(d))
            final = Document(resultado["docx"])
            processo = next(p for p in final.paragraphs if "Processo" in p.text)
            self.assertEqual(processo.style.name, "FB - Identificação do Processo")
            self.assertTrue(processo.style.font.bold)


if __name__ == "__main__":
    unittest.main()
