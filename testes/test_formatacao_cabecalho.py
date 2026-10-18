import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from scripts.aplicar_estilos import aplicar

ROOT = Path(__file__).resolve().parents[1]
MODELO = ROOT / "modelos/modelo_oficial.docx"
CONFIG = json.loads((ROOT / "modelos/estilos.json").read_text(encoding="utf-8"))
ESTILOS = ("FB - Endereçamento", "FB - Identificação do Processo")


class FormatacaoCabecalhoTests(unittest.TestCase):
    def test_modelo_e_configuracao_exigem_negrito_e_caixa_alta(self):
        doc = Document(MODELO)
        for nome in ESTILOS:
            fonte = doc.styles[nome].font
            self.assertTrue(fonte.bold)
            self.assertTrue(fonte.cs_bold)
            self.assertTrue(fonte.all_caps)
            self.assertFalse(fonte.small_caps)
            self.assertTrue(CONFIG["formatacao_estilos"][nome]["caixa_alta"])

    def test_formato_da_origem_nao_anula_regra_nem_altera_texto_e_campos(self):
        doc = Document()
        doc.add_paragraph("Excelentíssimo Senhor Desembargador Presidente do Tribunal de Justiça do Estado de Sergipe")
        doc.add_paragraph("Processo nº 202600000001 / 0001234-56.2025.8.26.0001")
        doc.add_paragraph("Texto do corpo em caixa mista, com número 123 e sigla TJSE.")
        link = OxmlElement("w:hyperlink")
        run = OxmlElement("w:r")
        texto = OxmlElement("w:t")
        texto.text = " complemento"
        run.append(texto)
        link.append(run)
        doc.paragraphs[0]._p.append(link)
        campo = OxmlElement("w:fldSimple")
        campo.set(qn("w:instr"), ' DOCVARIABLE "codigoProcesso" ')
        run = OxmlElement("w:r")
        texto = OxmlElement("w:t")
        texto.text = " CódigoAbC123"
        run.append(texto)
        campo.append(run)
        doc.paragraphs[1]._p.append(campo)
        for paragrafo in doc.paragraphs[:2]:
            for run in paragrafo._p.xpath(".//w:r"):
                rpr = run.get_or_add_rPr()
                for tag, val in (("b", "0"), ("bCs", "0"), ("caps", "0"), ("smallCaps", "1")):
                    flag = OxmlElement("w:" + tag)
                    flag.set(qn("w:val"), val)
                    rpr.append(flag)
        textos = [p._p.xpath(".//w:t/text()") for p in doc.paragraphs]
        with tempfile.TemporaryDirectory() as temp:
            entrada = Path(temp) / "entrada.docx"
            saida = Path(temp) / "saida.docx"
            doc.save(entrada)
            original = entrada.read_bytes()
            modelo_original = MODELO.read_bytes()
            classificacao = {"classificacoes": [
                {"id": 1, "estilo_destino": ESTILOS[0], "confianca": 1},
                {"id": 2, "estilo_destino": ESTILOS[1], "confianca": 1},
                {"id": 3, "estilo_destino": "FB - Corpo", "confianca": 1},
            ]}
            auditoria = aplicar(entrada, MODELO, classificacao, CONFIG, saida)
            self.assertEqual(auditoria["cabecalhos_negrito_caixa_alta"], 2)
            resultado = Document(saida)
            self.assertEqual([p._p.xpath(".//w:t/text()") for p in resultado.paragraphs], textos)
            self.assertEqual(resultado.paragraphs[1]._p.xpath(".//w:fldSimple/@w:instr"), [' DOCVARIABLE "codigoProcesso" '])
            for p in resultado.paragraphs[:2]:
                for rpr in [p._p.pPr.find(qn("w:rPr")), *p._p.xpath(".//w:r/w:rPr")]:
                    for tag in ("b", "bCs", "caps"):
                        flag = rpr.find(qn("w:" + tag))
                        self.assertIsNotNone(flag)
                        self.assertNotIn(flag.get(qn("w:val")), ("0", "off", "false"))
                    self.assertEqual(rpr.find(qn("w:smallCaps")).get(qn("w:val")), "0")
            corpo = resultado.paragraphs[2]
            self.assertFalse(corpo.style.font.bold)
            self.assertNotEqual(corpo.style.font.all_caps, True)
            segunda = Path(temp) / "segunda.docx"
            aplicar(saida, MODELO, classificacao, CONFIG, segunda)
            self.assertEqual([p._p.xpath(".//w:t/text()") for p in Document(segunda).paragraphs], textos)
            self.assertEqual(entrada.read_bytes(), original)
            self.assertEqual(MODELO.read_bytes(), modelo_original)
