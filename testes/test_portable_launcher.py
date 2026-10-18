import json
import tempfile
import unittest
from pathlib import Path

from docx import Document

from portable_launcher import (
    Aplicativo,
    carregar_dados_revisao,
    interface_tk_funcional,
    nome_seguro,
    processar_arquivo,
    reprocessar_resultado,
)
from tkinter import ttk


ROOT = Path(__file__).resolve().parents[1]


class PortableLauncherTests(unittest.TestCase):
    def test_nome_seguro_remove_caracteres_invalidos(self):
        self.assertEqual(nome_seguro('Petição: autor/réu?.docx'), "Petição_ autor_réu_")
        self.assertEqual(nome_seguro('Petição: autor/réu?.md'), "Petição_ autor_réu_")

    def test_paineis_internos_da_revisao_nao_possuem_borda(self):
        if not interface_tk_funcional():
            self.skipTest("Tk não está disponível neste ambiente.")
        app = Aplicativo(auto_iniciar=False)
        app.withdraw()
        try:
            estilo = ttk.Style(app)
            self.assertEqual(int(estilo.lookup("CardBody.TFrame", "borderwidth")), 0)
            self.assertGreaterEqual(int(estilo.lookup("Card.TFrame", "borderwidth")), 1)
        finally:
            app.destroy()

    def test_acoes_da_lista_ficam_abaixo_do_quadro(self):
        if not interface_tk_funcional():
            self.skipTest("Tk não está disponível neste ambiente.")
        app = Aplicativo(auto_iniciar=False)
        app.withdraw()
        try:
            self.assertGreater(
                int(app.acoes_lista.grid_info()["row"]),
                int(app.quadro_lista.grid_info()["row"]),
            )
            self.assertEqual(app.acoes_lista.grid_info()["sticky"], "e")
            self.assertIs(app.botao_limpar_lista.master, app.acoes_lista)
            self.assertIs(app.botao_remover_selecionados.master, app.acoes_lista)
        finally:
            app.destroy()

    def test_processamento_portatil_preserva_original_e_publica_resultados(self):
        with tempfile.TemporaryDirectory() as temporario:
            raiz = Path(temporario)
            entrada = raiz / "Petição de teste.docx"
            documento = Document()
            documento.add_paragraph("EXCELENTÍSSIMO SENHOR DOUTOR JUIZ DE DIREITO")
            documento.add_paragraph("Processo nº 0000000-00.2026.8.00.0000")
            documento.add_paragraph("Este trecho exige contexto para que sua função estrutural seja confirmada.")
            documento.add_paragraph("Diante do exposto, requer o regular prosseguimento.")
            documento.save(entrada)
            original = entrada.read_bytes()

            resultado = processar_arquivo(
                entrada,
                raiz / "resultados",
                gerar_pdf=False,
                modelos_raiz=ROOT / "modelos",
                execucoes_raiz=raiz / "execucoes_locais",
            )

            self.assertEqual(entrada.read_bytes(), original)
            self.assertTrue(resultado.docx.is_file())
            self.assertTrue(resultado.relatorio.is_file())
            self.assertIsNone(resultado.pdf)
            self.assertTrue((resultado.pasta_resultado / "auditoria.json").is_file())
            self.assertTrue((resultado.pasta_resultado / "resultado_execucao.json").is_file())
            resumo = json.loads(
                (resultado.pasta_resultado / "resultado_execucao.json").read_text(encoding="utf-8")
            )
            self.assertEqual(resumo["tipo_peca"], "peticao_inicial")
            self.assertTrue(resultado.execucao_local.is_dir())

            dados = carregar_dados_revisao(resultado)
            self.assertEqual(len(dados["itens"]), resultado.ambiguos)
            self.assertGreater(resultado.ambiguos, 0)
            self.assertFalse(resultado.aprovado)

            with self.assertRaisesRegex(ValueError, "Confirme todas as pendências"):
                reprocessar_resultado(resultado, [])

            revisado = reprocessar_resultado(resultado, [
                {
                    "id": item["id"],
                    "hash": item["hash"],
                    "estilo_destino": item["pre_classificacao"],
                    "justificativa": "Confirmado pela interface no teste.",
                }
                for item in dados["itens"]
            ])
            self.assertEqual(revisado.pasta_resultado, resultado.pasta_resultado)
            self.assertEqual(revisado.ambiguos, 0)
            self.assertTrue(revisado.aprovado)
            self.assertFalse((revisado.pasta_resultado / "pendencias_semanticas.json").exists())
            resumo_revisado = json.loads(
                (revisado.pasta_resultado / "resultado_execucao.json").read_text(encoding="utf-8")
            )
            self.assertEqual(resumo_revisado["pendencias_semanticas"], 0)
            self.assertTrue(resumo_revisado["aprovado"])

    def test_markdown_publica_manifesto_e_reutiliza_o_mesmo_intermediario_na_revisao(self):
        with tempfile.TemporaryDirectory() as temporario:
            raiz = Path(temporario)
            entrada = raiz / "Petição de teste.md"
            entrada.write_text(
                "## DOS FATOS\n\n"
                "Este trecho exige contexto para que sua função estrutural seja confirmada pela revisão humana.\n",
                encoding="utf-8",
            )
            original = entrada.read_bytes()

            resultado = processar_arquivo(
                entrada,
                raiz / "resultados",
                gerar_pdf=False,
                modelos_raiz=ROOT / "modelos",
                execucoes_raiz=raiz / "execucoes_locais",
            )
            self.assertEqual(entrada.read_bytes(), original)
            self.assertEqual(resultado.formato_original, "markdown")
            self.assertTrue(resultado.entrada_pipeline.is_file())
            self.assertEqual(resultado.entrada_pipeline, resultado.docx_intermediario)
            self.assertTrue((resultado.pasta_resultado / "entrada_intermediaria.docx").is_file())
            self.assertTrue((resultado.pasta_resultado / "manifesto_entrada.json").is_file())
            intermediario_antes = resultado.entrada_pipeline.read_bytes()

            dados = carregar_dados_revisao(resultado)
            revisado = reprocessar_resultado(resultado, [
                {
                    "id": item["id"],
                    "hash": item["hash"],
                    "estilo_destino": item["pre_classificacao"],
                    "justificativa": "Confirmado pela interface no teste Markdown.",
                }
                for item in dados["itens"]
            ])
            self.assertEqual(revisado.entrada_pipeline, resultado.entrada_pipeline)
            self.assertEqual(revisado.entrada_pipeline.read_bytes(), intermediario_antes)
            resumo = json.loads(
                (revisado.pasta_resultado / "resultado_execucao.json").read_text(encoding="utf-8")
            )
            self.assertEqual(resumo["formato_original"], "markdown")
            self.assertEqual(resumo["pendencias_semanticas"], 0)


if __name__ == "__main__":
    unittest.main()
