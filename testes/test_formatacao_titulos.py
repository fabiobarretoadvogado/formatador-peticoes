import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm

from scripts.aplicar_estilos import aplicar

ROOT = Path(__file__).resolve().parents[1]
MODELO = ROOT / "modelos/modelo_oficial.docx"
CONFIG = json.loads((ROOT / "modelos/estilos.json").read_text(encoding="utf-8"))
TITULOS = ("FB - Título Principal", "FB - Título Secundário")


class FormatacaoTitulosTests(unittest.TestCase):
    def conferir(self, doc):
        for nome in TITULOS:
            self.assertTrue(doc.styles[nome].font.bold)
            self.assertTrue(doc.styles[nome].font.cs_bold)
        estilo = doc.styles[TITULOS[1]]
        self.assertAlmostEqual(estilo.paragraph_format.left_indent.cm, 1, places=2)
        self.assertAlmostEqual(estilo.paragraph_format.first_line_indent.cm, 1, places=2)
        ids = {doc.styles[nome].style_id: nome for nome in TITULOS}
        encontrados = set()
        for nivel in doc.part.numbering_part.element.findall(".//" + qn("w:lvl")):
            vinculo = nivel.find(qn("w:pStyle"))
            if vinculo is None or vinculo.get(qn("w:val")) not in ids:
                continue
            nome = ids[vinculo.get(qn("w:val"))]
            encontrados.add(nome)
            for atributo in ("b", "bCs"):
                flag = nivel.find(qn("w:rPr")).find(qn("w:" + atributo))
                self.assertIsNotNone(flag)
                self.assertNotIn(flag.get(qn("w:val")), ("0", "false", "off"))
            if nome == TITULOS[1]:
                self.assertEqual(nivel.find(qn("w:suff")).get(qn("w:val")), "space")
                ind = nivel.find(qn("w:pPr")).find(qn("w:ind"))
                self.assertEqual(ind.get(qn("w:left")), "567")
                self.assertEqual(ind.get(qn("w:firstLine")), "567")
                for atributo in ("hanging", "hangingChars", "firstLineChars"):
                    self.assertIsNone(ind.get(qn("w:" + atributo)))
        self.assertEqual(encontrados, set(TITULOS))

    def test_modelo_oficial_e_catalogo_concordam_com_a_regra(self):
        self.conferir(Document(MODELO))
        self.assertEqual(CONFIG["formatacao_estilos"][TITULOS[1]]["recuo_primeira_linha_cm"], 1)

    def test_origem_sem_negrito_e_com_deslocamento_nao_prevalece(self):
        doc = Document()
        doc.add_paragraph("Título principal", style="Heading 1")
        p = doc.add_paragraph("Subtítulo extenso para conferir a primeira linha e a continuação do texto", style="Heading 2")
        p.paragraph_format.left_indent = Cm(1)
        p.paragraph_format.first_line_indent = Cm(-1)
        link = OxmlElement("w:hyperlink")
        run = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        b = OxmlElement("w:b")
        b.set(qn("w:val"), "0")
        rpr.append(b)
        run.append(rpr)
        texto = OxmlElement("w:t")
        texto.text = " complemento"
        run.append(texto)
        link.append(run)
        p._p.append(link)
        doc.add_paragraph("Texto do corpo que deve permanecer sem negrito.")
        for paragrafo in doc.paragraphs[:2]:
            for trecho in paragrafo.runs:
                trecho.bold = False
                trecho.font.cs_bold = False
            marca = OxmlElement("w:rPr")
            b = OxmlElement("w:b")
            b.set(qn("w:val"), "0")
            marca.append(b)
            paragrafo._p.get_or_add_pPr().append(marca)
        with tempfile.TemporaryDirectory() as temp:
            entrada, saida = Path(temp) / "entrada.docx", Path(temp) / "saida.docx"
            doc.save(entrada)
            original = entrada.read_bytes()
            modelo_original = MODELO.read_bytes()
            classificacao = {"classificacoes": [
                {"id": 1, "estilo_destino": TITULOS[0], "confianca": 1},
                {"id": 2, "estilo_destino": TITULOS[1], "confianca": 1},
                {"id": 3, "estilo_destino": "FB - Corpo", "confianca": 1},
            ]}
            aplicar(entrada, MODELO, classificacao, CONFIG, saida)
            self.assertEqual(entrada.read_bytes(), original)
            self.assertEqual(MODELO.read_bytes(), modelo_original)
            resultado = Document(saida)
            self.conferir(resultado)
            for indice, paragrafo in enumerate(resultado.paragraphs[:2]):
                self.assertEqual(paragrafo.style.name, TITULOS[indice])
                for rpr in [paragrafo._p.pPr.find(qn("w:rPr")), *paragrafo._p.xpath(".//w:r/w:rPr")]:
                    for atributo in ("b", "bCs"):
                        flag = rpr.find(qn("w:" + atributo))
                        self.assertIsNotNone(flag)
                        self.assertNotIn(flag.get(qn("w:val")), ("0", "false", "off"))
            subtitulo = resultado.paragraphs[1]
            self.assertAlmostEqual(subtitulo.paragraph_format.first_line_indent.cm, 1, places=2)
            self.assertIsNone(subtitulo._p.pPr.ind.get(qn("w:hanging")))
            self.assertIn(" complemento", "".join(subtitulo._p.xpath(".//w:t/text()")))
            self.assertEqual(resultado.paragraphs[2].text, doc.paragraphs[2].text)
            self.assertFalse(resultado.styles["FB - Corpo"].font.bold)
