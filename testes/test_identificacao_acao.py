import json
import tempfile
import unittest
from pathlib import Path

from scripts.classificar_regras import classificar, classificar_um


NOME_PECA = "CONTRARRAZÕES AO RECURSO ESPECIAL"


class IdentificacaoAcaoTests(unittest.TestCase):
    def paragrafo(self, texto, **extras):
        return {"id": 1, "texto": texto, "posicao_relativa": 0.10, **extras}

    def test_denominacao_da_imagem_e_identificacao_sem_revisao(self):
        item = classificar_um(
            self.paragrafo(NOME_PECA),
            "A parte autora, já qualificada, apresenta as presentes contrarrazões.",
            "EGRÉGIO SUPERIOR TRIBUNAL DE JUSTIÇA",
        )
        self.assertEqual(item["estilo_destino"], "FB - Identificação da Ação")
        self.assertEqual(item["confianca"], 0.99)
        self.assertFalse(item["necessita_analise_semantica"])
        self.assertEqual(item["texto"], NOME_PECA)

    def test_denominacao_admite_caixa_acentos_e_espacos_variados(self):
        for texto in (
            "Contrarrazões ao recurso especial",
            "CONTRARRAZOES AO RECURSO ESPECIAL",
            "Contrarrazões  ao\trecurso especial",
            "CONTRARRAZÕES AO RECURSO EXTRAORDINÁRIO",
            "Contrarrazoes ao recurso extraordinario",
        ):
            with self.subTest(texto=texto):
                item = classificar_um(self.paragrafo(texto))
                self.assertEqual(item["estilo_destino"], "FB - Identificação da Ação")

    def test_denominacao_isolada_prevalece_sobre_heading_original(self):
        for estilo in ("Heading 1", "Heading 2", "Título 1", "Título 2"):
            with self.subTest(estilo=estilo):
                item = classificar_um(self.paragrafo(NOME_PECA, estilo_original=estilo))
                self.assertEqual(item["estilo_destino"], "FB - Identificação da Ação")

    def test_mencoes_narrativas_e_secoes_nao_sao_nome_da_peca(self):
        for texto in (
            "AS CONTRARRAZÕES AO RECURSO ESPECIAL DEMONSTRAM O ERRO ALEGADO.",
            "Apresenta as presentes contrarrazões ao recurso especial.",
            "CONTRARRAZÕES AO RECURSO ESPECIAL NÃO FORAM APRESENTADAS",
            "DA ANÁLISE DAS CONTRARRAZÕES AO RECURSO ESPECIAL",
            "I. CONTRARRAZÕES AO RECURSO ESPECIAL",
        ):
            with self.subTest(texto=texto):
                item = classificar_um(self.paragrafo(texto))
                self.assertNotEqual(item["estilo_destino"], "FB - Identificação da Ação")
        for estilo in ("Quote", "Citação", "Citacao"):
            with self.subTest(estilo=estilo):
                item = classificar_um(self.paragrafo(NOME_PECA, estilo_original=estilo))
                self.assertEqual(item["estilo_destino"], "FB - Citação")

    def test_identificacao_de_acao_anterior_e_preservada(self):
        item = classificar_um(self.paragrafo("AÇÃO DE CUMPRIMENTO DE SENTENÇA"))
        self.assertEqual(item["estilo_destino"], "FB - Identificação da Ação")

    def test_docx_aplica_identificacao_sem_numeracao_de_secao(self):
        from docx import Document
        from docx.oxml.ns import qn
        from scripts.aplicar_estilos import aplicar
        from scripts.extrair_documento import extrair

        root = Path(__file__).resolve().parents[1]
        with tempfile.TemporaryDirectory() as temp:
            entrada, saida = Path(temp) / "entrada.docx", Path(temp) / "saida.docx"
            doc = Document()
            doc.add_paragraph("A parte autora, já qualificada, apresenta as presentes contrarrazões.")
            doc.add_paragraph(NOME_PECA, style="Heading 1")
            doc.add_paragraph("EGRÉGIO SUPERIOR TRIBUNAL DE JUSTIÇA")
            doc.save(entrada)
            original = entrada.read_bytes()
            config = json.loads((root / "modelos/estilos.json").read_text(encoding="utf-8"))
            regras = classificar(extrair(entrada))
            aplicar(entrada, root / "modelos/modelo_oficial.docx", regras, config, saida)
            p = Document(saida).paragraphs[1]
            self.assertEqual(p.style.name, "FB - Identificação da Ação")
            self.assertEqual(p.text, NOME_PECA)
            self.assertIsNone(p._p.pPr.find(qn("w:numPr")))
            self.assertEqual(entrada.read_bytes(), original)
